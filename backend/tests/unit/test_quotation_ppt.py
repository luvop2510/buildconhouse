"""The exported deck must preserve grouping and contain product imagery only."""
from io import BytesIO

from PIL import Image
from pptx import Presentation
import pytest

from models import QuotationLineItem
from services import quotation_ppt


def _image_bytes() -> bytes:
    stream = BytesIO()
    Image.new("RGB", (1200, 1200), "white").save(stream, format="PNG")
    return stream.getvalue()


def test_area_is_a_persisted_quotation_line_field():
    line = QuotationLineItem(product_id="p1", sku="S-1", name="Basin", room="Powder Room", area="Basin Area")
    assert line.model_dump()["area"] == "Basin Area"


def test_dimension_options_survive_quotation_line_validation():
    line = QuotationLineItem(product_id="p1", sku="S-1", name="Basin", dimension="600 × 400 mm", exclude_dimension=True)
    assert line.model_dump()["dimension"] == "600 × 400 mm"
    assert line.model_dump()["exclude_dimension"] is True


def test_ppt_groups_products_by_room_and_area_without_catalog_text(monkeypatch):
    data = _image_bytes()
    monkeypatch.setattr(quotation_ppt, "prefetch_urls", lambda _urls: None)
    monkeypatch.setattr(quotation_ppt, "remote_image_bytes", lambda _url: data)
    doc = {
        "customer_name": "Dineshbhai Amrutiya",
        "referrer_name": "Uday Garala",
        "referrer_type": "interior_designer",
        "rooms": ["Powder Room", "Utility"],
        "items": [
            {"room": "Powder Room", "area": "Basin Area", "name": "Basin", "sku": "SECRET-1", "image": "https://example.test/1.png"},
            {"room": "Powder Room", "area": "Basin Area", "name": "Tap", "sku": "SECRET-2", "image": "https://example.test/2.png"},
            {"room": "Powder Room", "area": "WC Area", "name": "WC", "sku": "SECRET-3", "image": "https://example.test/3.png"},
            {"room": "Utility", "area": "Shower Area", "name": "Shower", "sku": "SECRET-4", "image": "https://example.test/4.png"},
        ],
    }
    ppt = Presentation(BytesIO(quotation_ppt.build_product_ppt(doc)))
    names = {row[0] for row in quotation_ppt.PARTNER_LOGOS}
    assert "hansgrohe" in names and "axor" not in names
    assert len(ppt.slides) == 3
    assert round(ppt.slide_width / ppt.slide_height, 2) == round(16 / 9, 2)
    texts = [" ".join(shape.text for shape in slide.shapes if shape.has_text_frame) for slide in ppt.slides]
    assert "CLIENT NAME: DINESHBHAI AMRUTIYA" in texts[0]
    assert "INT: UDAY GARALA" in texts[0]
    assert "9909906652" in texts[0]
    assert "POWDER ROOM" in texts[1] and "Basin Area" in texts[1] and "WC Area" in texts[1]
    assert "UTILITY" in texts[2] and "Shower Area" in texts[2]
    assert all("SECRET" not in text and "Tap" not in text for text in texts)
    assert [sum(shape.shape_type == 13 for shape in slide.shapes) for slide in ppt.slides] == [10, 3, 1]


def test_cover_uses_architect_label_when_referrer_is_architect(monkeypatch):
    monkeypatch.setattr(quotation_ppt, "prefetch_urls", lambda _urls: None)
    ppt = Presentation(BytesIO(quotation_ppt.build_product_ppt({
        "customer_name": "Another Client", "referrer_name": "Patel Studio",
        "referrer_type": "architect", "items": [{"name": "Basin", "room": "Bath"}],
    })))
    cover_text = " ".join(shape.text for shape in ppt.slides[0].shapes if shape.has_text_frame)
    assert "CLIENT NAME: ANOTHER CLIENT" in cover_text
    assert "ARCH: PATEL STUDIO" in cover_text
    assert "DINESHBHAI" not in cover_text


def test_ppt_keeps_product_with_missing_image():
    ppt = Presentation(BytesIO(quotation_ppt.build_product_ppt({"items": [{"name": "Basin", "room": "Bath"}]})))
    assert len(ppt.slides) == 2
    assert sum(shape.shape_type == 13 for shape in ppt.slides[1].shapes) == 1


def test_filename_contains_only_safe_customer_name_ppt_and_date():
    assert quotation_ppt.presentation_filename('A/B: Customer', "25-09-2026") == "AB Customer PPT 25-09-2026.pptx"


def test_unassigned_legacy_products_get_sensible_areas():
    assert quotation_ppt.suggested_area("Wall hung WC with bidet") == "WC Area"
    assert quotation_ppt.suggested_area("Rainshower 300") == "Shower Area"
    assert quotation_ppt.suggested_area("Countertop basin") == "Basin Area"
    assert quotation_ppt.suggested_area("Towel rail") == "Other Area"


@pytest.mark.parametrize("count,area_count", [(7, 1), (17, 3), (36, 4), (60, 6)])
def test_every_room_stays_on_one_slide_with_all_images_inside_nonoverlapping_frames(monkeypatch, count, area_count):
    monkeypatch.setattr(quotation_ppt, "prefetch_urls", lambda _urls: None)
    monkeypatch.setattr(quotation_ppt, "remote_image_bytes", lambda _url: _image_bytes())
    items = [{"room": "Master Bathroom", "area": f"Area {i % area_count}", "image": f"https://test/{i}.png"}
             for i in range(count)]
    items.append({"room": "Powder Room", "area": "WC Area", "image": "https://test/wc.png"})
    ppt = Presentation(BytesIO(quotation_ppt.build_product_ppt({"items": items})))
    assert len(ppt.slides) == 3
    photos = [shape for shape in ppt.slides[1].shapes if shape.shape_type == 13]
    assert len(photos) == count
    for photo in photos:
        assert 0 < photo.left < photo.left + photo.width <= ppt.slide_width
        assert 0 < photo.top < photo.top + photo.height <= ppt.slide_height
        assert photo.crop_left == photo.crop_right == photo.crop_top == photo.crop_bottom == 0
    for index, first in enumerate(photos):
        for second in photos[index + 1:]:
            assert (first.left + first.width <= second.left or second.left + second.width <= first.left
                    or first.top + first.height <= second.top or second.top + second.height <= first.top)


def test_normalized_room_labels_and_duplicate_declared_rooms_never_lose_or_duplicate_products(monkeypatch):
    monkeypatch.setattr(quotation_ppt, "prefetch_urls", lambda _urls: None)
    ppt = Presentation(BytesIO(quotation_ppt.build_product_ppt({
        "rooms": [" Bath ", "Bath", "Empty"],
        "items": [{"room": " Bath "}, {"room": "Bath"}, {"room": " "}, {}],
    })))
    assert len(ppt.slides) == 3
    assert [sum(shape.shape_type == 13 for shape in slide.shapes) for slide in list(ppt.slides)[1:]] == [2, 2]


def test_seat_and_concealed_parts_are_present_and_repeated_assets_only_normalized_once(monkeypatch):
    data = _image_bytes()
    monkeypatch.setattr(quotation_ppt, "prefetch_urls", lambda _urls: None)
    monkeypatch.setattr(quotation_ppt, "remote_image_bytes", lambda _url: data)
    normalized = []
    original = quotation_ppt._image

    def prepare(image):
        normalized.append(image)
        return original(image)

    monkeypatch.setattr(quotation_ppt, "_image", prepare)
    ppt = Presentation(BytesIO(quotation_ppt.build_product_ppt({"items": [
        {"room": "Bath", "name": "WC", "image": "https://test/wc.png", "seat_cover_image_urls": ["https://test/wc.png", "https://test/seat.png"]},
        {"room": "Bath", "name": "Concealed part", "image": "https://test/rough-in.png"},
        {"room": "Powder", "name": "WC", "image": "https://test/wc.png"},
    ]})))
    assert len(ppt.slides) == 3
    assert sum(shape.shape_type == 13 for shape in ppt.slides[1].shapes) == 3
    assert sum(shape.shape_type == 13 for shape in ppt.slides[2].shapes) == 1
    assert len(normalized) == 3


def test_large_portrait_product_is_resized_without_distortion_or_cropping():
    source = BytesIO()
    Image.new("RGB", (2000, 4000), "blue").save(source, format="PNG")
    stream, ratio = quotation_ppt._image(source.getvalue())
    with Image.open(stream) as image:
        assert image.size == (800, 1600)
        assert ratio == 0.5
