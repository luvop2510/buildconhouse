"""Assert emitted CSV/XLSX bytes keep untrusted cells as text."""
import csv
from io import BytesIO, StringIO
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from openpyxl import load_workbook
from models import UserPublic


class Collection:
    def __init__(self, rows):
        self.rows = rows

    def find(self, *args, **kwargs):
        return self

    def sort(self, *args):
        return self

    async def to_list(self, limit):
        return self.rows


async def response_bytes(response):
    chunks = [chunk async for chunk in response.body_iterator]
    return b"".join(chunk.encode() if isinstance(chunk, str) else chunk for chunk in chunks)


def owner():
    return UserPublic(id="owner", email="owner@example.com", full_name="Owner", role="owner")


@pytest.mark.asyncio
async def test_catalog_formula_text_and_numeric_prices(monkeypatch):
    from routes import catalog_routes as routes
    monkeypatch.setattr(routes, "db", SimpleNamespace(
        products=Collection([{"sku": "=1+1", "name": "\t@SUM(1)", "brand_id": "b", "category_id": "c", "mrp": 123.5}]),
        brands=Collection([{"id": "b", "name": "+1+1"}]), categories=Collection([{"id": "c", "name": "-1+1"}]),
    ))
    sheet = load_workbook(BytesIO(await response_bytes(await routes.export_catalog_xlsx(owner())))).active
    for cell in sheet[5][:4]:
        assert cell.data_type != "f"
        assert cell.value.startswith("'")
    assert sheet["G5"].value == 123.5
    assert sheet["G5"].data_type == "n"


@pytest.mark.asyncio
@pytest.mark.parametrize("format", ["csv", "xlsx"])
async def test_followup_formula_injection_in_emitted_file(monkeypatch, format):
    from routes import followup_routes as routes
    monkeypatch.setattr(routes, "list_followups", AsyncMock(return_value=[{
        "customer_name": "=1+1", "reason": " @SUM(1)", "value": 50,
    }]))
    payload = await response_bytes(await routes.export_followups(format=format, user=owner()))
    if format == "csv":
        row = list(csv.reader(StringIO(payload.decode())))[1]
        assert row[0] == "'=1+1"
        assert row[3].startswith("'")
        assert row[5] == "50"
    else:
        sheet = load_workbook(BytesIO(payload)).active
        assert sheet["A5"].data_type != "f"
        assert sheet["D5"].value.startswith("'")
        assert sheet["F5"].value == 50


@pytest.mark.asyncio
@pytest.mark.parametrize("format", ["csv", "xlsx"])
async def test_executive_formula_injection_in_emitted_file(monkeypatch, format):
    from routes import executive_analytics_routes as routes
    monkeypatch.setattr(routes, "db", SimpleNamespace(quotations=Collection([
        {"number": "=1+1", "customer_name": "+SUM(1)", "grand_total": 75},
    ])))
    payload = await response_bytes(await routes.export_orders(format=format, preset="all", user=owner()))
    if format == "csv":
        row = list(csv.reader(StringIO(payload.decode())))[1]
        assert row[0] == "'=1+1"
        assert row[3] == "'+SUM(1)"
        assert row[5] == "75"
    else:
        sheet = load_workbook(BytesIO(payload)).active
        assert sheet["A2"].data_type != "f"
        assert sheet["D2"].value == "'+SUM(1)"
        assert sheet["F2"].value == 75
