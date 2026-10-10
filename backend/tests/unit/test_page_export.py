from io import BytesIO

import openpyxl
import pytest
from fastapi import HTTPException
from models import UserPublic
from routes import page_export_routes

from routes.page_export_routes import ExportSection, PageWorkbook, export_page_workbook


@pytest.mark.asyncio
async def test_page_workbook_keeps_sections_and_escapes_spreadsheet_formulas():
    payload = PageWorkbook(page="Catalog", sections=[
        ExportSection(name="Brand A", rows=[["SKU", "Price"], ["=HYPERLINK(\"bad\")", "1200"]]),
        ExportSection(name="Brand B", rows=[["SKU", "Price"], ["B2", "900"]]),
    ])

    response = await export_page_workbook(payload, user=UserPublic(email="owner@example.com", full_name="Owner", role="owner"))
    data = b"".join([chunk async for chunk in response.body_iterator])
    workbook = openpyxl.load_workbook(BytesIO(data))

    assert workbook.sheetnames == ["Brand A", "Brand B 2"]
    assert workbook["Brand A"]["A4"].value.startswith("'=HYPERLINK")
    assert workbook["Brand B 2"]["A4"].value == "B2"


@pytest.mark.asyncio
@pytest.mark.parametrize("resource,floor,actions,allowed", [
    (None, "ground-floor", ["view", "export"], False),
    ("customers", None, ["view", "export"], False),
    ("customers", "ground-floor", ["view"], False),
    ("customers", "first-floor", ["view", "export"], False),
    ("customers", "ground-floor", ["view", "export"], True),
])
async def test_page_export_requires_exact_resource_floor_and_export(monkeypatch, resource, floor, actions, allowed):
    user = UserPublic(id="staff", email="staff@example.com", full_name="Staff", role="worker",
                      custom_access=True, floor_ids=["ground-floor"], active_floor_id=floor)
    async def grants(_database, _user_id):
        return [{"id": "grant", "user_id": "staff", "resource": "customers",
                 "floor_id": "ground-floor", "actions": actions}]
    monkeypatch.setattr(page_export_routes, "grants_for_user", grants)
    payload = PageWorkbook(page="Customers", resource=resource, sections=[ExportSection(name="Customers", rows=[["Name"], ["Example"]])])
    if allowed:
        response = await export_page_workbook(payload, user=user)
        assert response.status_code == 200
    else:
        with pytest.raises(HTTPException) as error:
            await export_page_workbook(payload, user=user)
        assert error.value.status_code == 403
