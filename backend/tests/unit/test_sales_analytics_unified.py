"""Financial invariants and authorized boundaries of the shared analytics model."""
import asyncio
from dataclasses import replace
from datetime import datetime, timezone
from io import BytesIO
import time

import pytest
from fastapi import HTTPException, Response
from openpyxl import load_workbook
from models import UserPublic
from routes import sales_analytics_routes as routes
from services.analytics import unified as sales
from services.analytics.unified_exports import export_bytes, report_sections
from services.analytics.unified_profiles import profile
from services.analytics.filters import FloorAccessError

NOW = datetime(2026, 10, 28, tzinfo=timezone.utc)


def order(identity="o1", customer="c1", amount=100, *, date="2026-10-10T12:00:00+00:00", floor="ground-floor", lines=None, payments=None):
    doc = {"id": identity, "number": identity.upper(), "floor_id": floor, "customer_id": customer, "customer_name": customer or "Unknown", "status": "ordered", "ordered_at": date, "created_at": "2026-09-01T10:00:00+00:00", "grand_total": amount, "referrer_id": "a1", "referrer_name": "Architect", "referrer_type": "architect", "items": lines if lines is not None else [{"id": identity + ":l", "product_id": "p1", "name": "Tile", "sku": "T1", "qty": 2, "unit_price": amount / 2 if amount else 0, "net_amount": amount, "quantity_unit": "Box", "size": "600x600", "finish": "Matt"}]}
    return sales.normalize("catalog", doc, {"p1": {"id": "p1", "brand_id": "b1"}, "p2": {"id": "p2", "brand_id": "b2"}}, {"b1": "Company One", "b2": "Company Two"}, payments or [])


def test_multiline_order_counts_once_and_brands_reconcile():
    lines = [{"product_id": "p1", "name": "Tile A", "qty": 2, "net_amount": 60}, {"product_id": "p2", "name": "Tile B", "qty": 1, "net_amount": 40}]
    report = sales.dashboard([order(lines=lines)], sales.Query(), None, NOW)
    assert report["kpis"]["sales"] == 100
    assert report["kpis"]["orders"] == report["kpis"]["customers"] == 1
    assert sum(r["sales"] for r in report["companies"]) == 100
    assert next(r for r in report["companies"] if r["id"] == "b1")["average_order_value"] == 60
    assert sum(t["sales"] for t in report["trend"]) == 100


def test_payment_basis_uses_cash_not_product_order_value_or_credits():
    row = order(amount=1000, payments=[{"id": "pay1", "amount": 100, "paid_at": "2026-10-01T00:00:00+00:00"}, {"id": "pay2", "amount": 250, "paid_at": "2026-10-15T00:00:00+00:00"}, {"id": "credit", "amount": 50, "paid_at": "2026-10-16T00:00:00+00:00", "entry_type": "godown_stock_credit"}])
    q = sales.Query(date_basis="payment", granularity="day")
    report = sales.dashboard([row], q, None, NOW)
    assert report["metric_label"] == "Collected amount"
    assert report["kpis"]["sales"] == report["kpis"]["collected"] == 350
    assert sum(r["sales"] for r in report["trend"]) == 350
    assert report["kpis"]["outstanding"] == 600
    selected = sales.select_period([row], replace(q, company_id="b1"), sales.period_for(q, NOW))
    assert selected[0]["sales"] is None and selected[0]["lines"][0]["sales"] is None


def test_currency_separation_and_aov_denominators():
    rnb = sales.normalize("furniture", {"id": "f1", "floor_id": "third-floor", "customer_id": "c1", "customer_name": "Customer", "status": "final", "total_payable_rmb": "500", "total_received_rmb": "100", "balance_rmb": "400", "finalized_at": "2026-10-09T10:00:00+00:00"}, {}, {}, [])
    report = sales.dashboard([order(amount=100), rnb], sales.Query(), None, NOW)
    assert report["kpis"]["sales"] == 100 and report["kpis"]["sales_rnb"] == 500
    assert report["kpis"]["orders"] == 2 and report["kpis"]["customers"] == 1
    assert report["kpis"]["average_order_value"] == 100 and report["kpis"]["average_order_value_rnb"] == 500
    assert all(r["contribution_pct"] in {None, 0, 100} for r in report["floors"])


def test_missing_values_are_unavailable_and_names_do_not_merge_ids():
    rows = [order(amount=None, customer=None), order("o2", "different-stable-id", 50)]
    rows[1]["customer_name"] = rows[0]["customer_name"]
    report = sales.dashboard(rows, sales.Query(), None, NOW)
    assert report["kpis"]["sales"] is None
    assert report["kpis"]["customers"] is None and report["kpis"]["known_customers"] == 1
    assert report["coverage"]["missing_customer_ids"] == 1 and len(report["customers"]) == 1


def test_missing_confirmation_never_falls_back_to_created_at():
    report = sales.dashboard([order(date=None)], sales.Query(), None, NOW)
    assert report["kpis"]["orders"] == 0 and report["coverage"]["missing_dates"] == 1
    quote = sales.dashboard([order(date=None)], sales.Query(date_basis="quotation", preset="all"), None, NOW)
    assert quote["kpis"]["orders"] == 1


def test_india_boundaries_aliases_and_exclusive_iso_end():
    q = sales.Query(preset="custom", date_from="2026-10-01", date_to="2026-10-01")
    p = sales.period_for(q, NOW)
    assert p.start == "2026-09-30T18:30:00+00:00" and p.end == "2026-10-01T18:30:00+00:00"
    assert sales.period_for(replace(q, date_from=p.start, date_to=p.end), NOW) == p
    assert sales.period_for(sales.Query(preset="this_quarter"), NOW).label == "This quarter"
    assert sales.period_for(sales.Query(preset="this_year"), NOW).label == "This year"
    assert sales.period_for(sales.Query(preset="all_time"), NOW).start is None
    with pytest.raises(HTTPException):
        sales.period_for(replace(q, date_from="invalid"), NOW)


def test_new_and_repeat_drilldown_records_match_segment_counts():
    rows = [order("old", "repeat", date="2026-08-01T00:00:00+00:00"), order("repeat", "repeat"), order("new", "new")]
    q = sales.Query()
    report = sales.dashboard(rows, q, None, NOW)
    assert [r["customers"] for r in report["customer_segments"]] == [1, 1]
    selected = sales.select_period(rows, replace(q, dimension="repeat"), sales.period_for(q, NOW))
    assert [r["id"] for r in selected] == ["repeat"]


def test_floor_rejection_precedes_all_database_reads(monkeypatch):
    async def forbidden(*args):
        raise AssertionError("Unauthorized request performed a read")
    monkeypatch.setattr(routes, "load_model", forbidden)
    user = UserPublic(id="restricted", email="restricted@example.test", full_name="Restricted", role="admin", floor_ids=["first-floor"])
    with pytest.raises(HTTPException) as error:
        asyncio.run(routes.get_dashboard(Response(), q=sales.Query(floor_id="ground-floor"), user=user))
    assert error.value.status_code == 403
    with pytest.raises(FloorAccessError):
        sales.scope_for(sales.Query(floor_id="second-floor"), ["first-floor"])
    assert sales.scope_for(sales.Query(), [])["floor_id"] == {"$in": []}


def test_all_routes_role_gate_and_export_grant_mapping():
    from services.access_grants import resource_for_api_path, action_for_api_request
    assert resource_for_api_path("/api/sales-analytics/entities/customers/c1") == "sales_data"
    assert action_for_api_request("GET", "/api/sales-analytics/export", {"format": "pdf"}) == "export"
    for route in routes.router.routes:
        deps = [d.call for d in route.dependant.dependencies if getattr(d.call, "__name__", "") == "_dep"]
        assert deps
        with pytest.raises(HTTPException) as error:
            asyncio.run(deps[0](UserPublic(id="sales", email="sales@example.test", full_name="Sales", role="sales")))
        assert error.value.status_code == 403


def test_sources_exclude_drafts_supplier_pos_and_cap_explicitly():
    captures = []
    class Cursor:
        async def to_list(self, length):
            return []
    class Collection:
        def aggregate(self, pipeline, **kwargs):
            captures.append(pipeline)
            return Cursor()
    class DB:
        quotations = kitchen_projects = furniture_quotations = Collection()
    assert asyncio.run(sales.load_model(DB(), sales.Query(), None)) == []
    assert captures[0][0]["$match"]["status"] == {"$in": ["ordered", "won"]}
    assert "_confirmed" in str(captures[1]) and "$_basis.order" in str(captures[1])
    assert captures[2][0]["$match"]["status"] == "final"
    assert "purchase_orders" not in str(captures)
    class LargeCursor:
        async def to_list(self, length):
            return [{}] * length
    with pytest.raises(HTTPException) as error:
        asyncio.run(sales._bounded(LargeCursor(), "Sales", limit=3))
    assert error.value.status_code == 422


def test_profile_exports_single_identity_and_source_ids_to_real_formats():
    class Cursor:
        async def to_list(self, length):
            return []
    class Collection:
        async def find_one(self, match, fields):
            return {"id": "b1", "name": "Company One"}
        def find(self, match, fields):
            return Cursor()
    class DB:
        quotations = kitchen_projects = activity_events = products = Collection()
        def __getitem__(self, key):
            return Collection()
    q = sales.entity_query(sales.Query(), "companies", "b1")
    report = asyncio.run(profile(DB(), [order("one", amount=60), order("two", amount=40)], q, None, "companies", "b1", table_page_size=100))
    assert report["identity"]["id"] == "b1"
    sections = report_sections(report, view="profile")
    assert {r["id"] for name, rows in sections if name == "Orders" for r in rows} == {"one", "two"}
    book = load_workbook(BytesIO(export_bytes(report, sections, "xlsx", {"Date basis": "confirmation", "Selected entity ID": "b1"})))
    assert "Orders" in book.sheetnames and "Identity" in book.sheetnames
    assert list(book["Orders"].values)[1][0] == "one"
    assert export_bytes(report, sections, "pdf", {"Selected entity ID": "b1"}).startswith(b"%PDF")


def test_partial_export_rejection_and_formula_escaping():
    with pytest.raises(HTTPException):
        report_sections({"title": "Report", "tables": {"orders": {"items": [], "total": 1}}}, view="profile")
    data = export_bytes({}, [("Customers", [{"id": '=HYPERLINK("x")', "name": "Safe", "amount": 10}])], "xlsx", {})
    sheet = load_workbook(BytesIO(data))["Customers"]
    assert sheet["A2"].value.startswith("'=") and sheet["C2"].value == 10


def test_current_overview_export_is_analysis_not_order_list():
    report = {"kpis": {"sales": 100}, "tables": {"orders": {"items": [{"id": "one"}], "total": 1}}, "trend": [{"sales": 100}], "floors": [], "companies": [], "architects": [], "products": [], "customers": []}
    names = [name for name, _ in report_sections(report, view="profile", scope="current", tab="overview")]
    assert "Trend" in names and "Orders" not in names


def test_quantity_units_are_not_added_together_and_series_is_filterable():
    one, two = order(), order("two")
    two["lines"][0]["unit"] = "Pieces"
    report = sales.dashboard([one, two], sales.Query(), None, NOW)
    assert report["trend"][0]["quantity"] is None and report["products"][0]["quantity"] is None
    assert report["attribute_breakdowns"]["size"][0]["orders"] == 2
    one["lines"][0]["series"] = "Alpha"
    assert sales.matches(one, sales.Query(series="Alpha"))
    assert not sales.matches(one, sales.Query(series="Beta"))


def test_business_scale_multiline_model_reconciles_without_unbounded_work():
    rows = [order(str(i), "c" + str(i % 100), 100) for i in range(2000)]
    started = time.monotonic()
    report = sales.dashboard(rows, sales.Query(granularity="day"), None, NOW)
    assert report["kpis"]["sales"] == 200000 and report["kpis"]["orders"] == 2000 and report["kpis"]["customers"] == 100
    assert sum(row["sales"] for row in report["trend"]) == report["kpis"]["sales"]
    assert time.monotonic() - started < 20


def test_actual_entity_directory_contains_zero_sale_catalog_and_paginates():
    from services.analytics.unified_profiles import directory
    captures = []
    class Cursor:
        async def to_list(self, length):
            return [{"id": "p1", "name": "Sold tile", "sku": "SKU1", "floor_id": "ground-floor"}, {"id": "p2", "name": "Unsold tile", "sku": "SKU2", "floor_id": "ground-floor"}]
    class Collection:
        def find(self, query, fields):
            captures.append(query)
            return Cursor()
    class DB:
        def __getitem__(self, key):
            assert key == "products"
            return Collection()
    q = sales.Query(floor_id="ground-floor")
    result = asyncio.run(directory(DB(), [order()], q, ["ground-floor"], "products", page=2, page_size=1))
    assert result["total"] == 2 and result["items"][0]["id"] == "p2" and result["items"][0]["sales"] == 0
    assert captures[0]["floor_id"] == {"$in": ["ground-floor"]}


def test_http_routes_enforce_role_floor_input_and_export_record_scope(monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from auth import get_current_user
    app = FastAPI()
    app.include_router(routes.router, prefix="/api")
    principal = {"user": UserPublic(id="admin", email="a@example.test", full_name="Admin", role="admin", floor_ids=["ground-floor"])}
    async def user():
        return principal["user"]
    app.dependency_overrides[get_current_user] = user
    async def fake_model(db, q, floors):
        rows = [order("one", "c1", 60), order("two", "c2", 40)]
        return [row for row in rows if not q.customer_id or row["customer_id"] == q.customer_id]
    monkeypatch.setattr(routes, "load_model", fake_model)
    with TestClient(app) as client:
        assert client.get("/api/sales-analytics/dashboard?floor_id=first-floor").status_code == 403
        assert client.get("/api/sales-analytics/dashboard?date_basis=invalid").status_code == 422
        response = client.get("/api/sales-analytics/records?floor_id=ground-floor&preset=all&customer_id=c1")
        assert response.status_code == 200 and response.headers["cache-control"] == "no-store"
        assert [row["id"] for row in response.json()["items"]] == ["one"]
        export = client.get("/api/sales-analytics/export?format=xlsx&view=records&floor_id=ground-floor&preset=all&customer_id=c1")
        assert export.status_code == 200
        book = load_workbook(BytesIO(export.content))
        assert [row[0] for row in list(book["Records"].values)[1:]] == ["one"]
        principal["user"] = UserPublic(id="sales", email="s@example.test", full_name="Sales", role="sales")
        assert client.get("/api/sales-analytics/dashboard").status_code == 403
        assert client.get("/api/sales-analytics/export?format=pdf").status_code == 403


def test_custom_profile_redacts_contact_fields_in_identity_read():
    from services.analytics.unified_profiles import identity_for
    captured = []
    class Collection:
        async def find_one(self, match, fields):
            captured.append(fields)
            return {"id": "c1", "name": "Customer"}
    class DB:
        def __getitem__(self, key):
            return Collection()
    identity = asyncio.run(identity_for(DB(), sales.Query(), None, "customers", "c1", [], contacts=False))
    assert not {"phone", "email", "address"}.intersection(captured[0])
    assert not any(a["label"] in {"Phone", "Email", "Address"} for a in identity["attributes"])


def test_india_chart_buckets_exactly_partition_selected_midnight_orders():
    q = sales.Query(preset="custom", date_from="2026-10-01", date_to="2026-10-02", granularity="day")
    rows = [order("before", amount=1, date="2026-09-30T18:29:59+00:00"), order("oct1", amount=20, date="2026-09-30T18:30:00+00:00"), order("oct1-end", amount=30, date="2026-10-01T18:29:59+00:00"), order("oct2", amount=50, date="2026-10-01T18:30:00+00:00"), order("after", amount=2, date="2026-10-02T18:30:00+00:00")]
    report = sales.dashboard(rows, q, None, NOW)
    assert [point["label"] for point in report["trend"]] == ["01 Oct", "02 Oct"]
    assert [point["sales"] for point in report["trend"]] == [50, 50]
    assert report["kpis"]["sales"] == sum(point["sales"] for point in report["trend"]) == 100
    for point in report["trend"]:
        drill = replace(q, date_from=point["start"], date_to=point["end"])
        assert sales.totals(sales.select_period(rows, drill, sales.period_for(drill, NOW)))["sales"] == point["sales"]


def test_confirmation_extra_costs_use_append_only_history_over_stale_cache():
    row = sales.normalize("catalog", {"id": "o", "floor_id": "ground-floor", "grand_total": 100, "payment_extra_amount": 99, "payment_extra_history": [{"amount": 10}, {"amount": 5}]}, {}, {}, [{"id": "p", "amount": 20}])
    assert row["outstanding"] == 95


def test_profile_has_actual_draft_quotation_history_without_adding_sales():
    from services.analytics.unified_profiles import profile_history
    class Cursor:
        def __init__(self, rows):
            self.rows = rows
        async def to_list(self, length):
            return self.rows
    class Collection:
        def __init__(self, rows):
            self.rows = rows
            self.query = None
        def find(self, match, fields):
            self.query = match
            return Cursor(self.rows)
    class DB:
        quotations = Collection([{"id": "draft", "number": "Q-DRAFT", "floor_id": "ground-floor", "customer_id": "c1", "customer_name": "Customer", "created_at": "2026-10-02T10:00:00+00:00", "status": "draft", "grand_total": 2000}])
        kitchen_projects = Collection([])
        activity_events = Collection([{"id": "event", "quotation_id": "draft", "created_at": "2026-10-02T10:00:00+00:00", "event_type": "quotation.created"}])
    database = DB()
    q = sales.Query(floor_id="ground-floor", customer_id="c1")
    history = asyncio.run(profile_history(database, q, ["ground-floor"]))
    assert history["quotations"][0]["id"] == "draft" and history["quotations"][0]["quotation_amount"] == 2000
    assert history["quotations"][0]["order_id"] is None
    assert history["activity"][0]["id"] == "event"
    assert database.quotations.query["customer_id"] == "c1"
    assert "status" not in database.quotations.query
    report = sales.dashboard([], q, ["ground-floor"], NOW)
    assert report["kpis"]["sales"] == 0 and report["kpis"]["orders"] == 0


def test_fractional_quantities_preserve_four_decimals_and_invalid_money_is_unavailable():
    row = order(lines=[{"product_id": "p1", "name": "Tile", "qty": 1.2345, "net_amount": 100}])
    assert row["lines"][0]["quantity"] == 1.2345
    report = sales.dashboard([row], sales.Query(), None, NOW)
    assert report["products"][0]["quantity"] == 1.2345 and report["trend"][0]["quantity"] == 1.2345
    invalid = order(payments=[{"id": "invalid", "amount": "invalid", "paid_at": "2026-10-01T00:00:00+00:00"}])
    assert invalid["collected"] is None and invalid["outstanding"] is None
    data = sales.dashboard([invalid], sales.Query(), None, NOW)
    assert data["kpis"]["sales"] == 100 and data["kpis"]["collected"] is None
    corrupted_extra = sales.normalize("catalog", {"id": "o", "floor_id": "ground-floor", "grand_total": 100, "payment_extra_history": [{"amount": "bad"}]}, {}, {}, [])
    assert corrupted_extra["outstanding"] is None


def test_product_rankings_do_not_allocate_whole_order_collections_or_unknown_quantity():
    row = order(lines=[{"product_id": "p1", "name": "Tile", "sku": "Recorded-SKU", "qty": 1.2345, "net_amount": 60}, {"product_id": "p1", "name": "Tile", "qty": None, "net_amount": 40}], payments=[{"id": "p", "amount": 30}])
    product = sales.ranking([row], "products")[0]
    assert product["sales"] == 100 and product["quantity"] is None
    assert product["collected"] is None and product["outstanding"] is None
    assert product["sku"] == "Recorded-SKU"


def test_furniture_recorded_charges_are_lines_without_catalog_brand_identity():
    row = sales.normalize("furniture", {"id": "f1", "floor_id": "third-floor", "status": "final", "finalized_at": "2026-10-10T00:00:00+00:00", "customer_id": "c1", "total_payable_rmb": "25", "total_received_rmb": "5", "balance_rmb": "20", "charges": [{"id": "charge1", "description": "Sofa", "amount_rmb": "25"}]}, {}, {}, [])
    report = sales.dashboard([row], sales.Query(floor_id="third-floor"), None, NOW)
    assert report["kpis"]["sales_rnb"] == 25 and not report["companies"] and not report["products"]
    assert row["lines"][0]["name"] == "Sofa"
