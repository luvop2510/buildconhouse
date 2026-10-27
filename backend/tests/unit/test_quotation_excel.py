"""The Excel reference is a layout template, never a source of customer data."""
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
import pytest

from services import quotation_excel as excel


def _line(**overrides):
    return {"id": "l1", "product_id": "p1", "sku": "000123", "name": "Basin mixer", "unit_price": 800,
        "mrp": 1000, "qty": 2, "room": "Master", **overrides}


def _book(doc):
    return excel.build_quotation_excel(doc, {"name": "Customer", "phone": "1234567890"})


def test_template_static_parts_and_layout_are_preserved():
    data = _book({"items": [_line()], "created_at": "2026-10-26T00:00:00Z"})
    with ZipFile(excel.TEMPLATE) as reference, ZipFile(BytesIO(data)) as actual:
        for part in ("xl/styles.xml", "xl/theme/theme1.xml", "xl/drawings/drawing1.xml", "xl/media/image1.png", "xl/printerSettings/printerSettings1.bin"):
            assert actual.read(part) == reference.read(part)
        ns = {"s": excel.S}
        for index in (1, 2):
            old = ET.fromstring(reference.read(f"xl/worksheets/sheet{index}.xml"))
            new = ET.fromstring(actual.read(f"xl/worksheets/sheet{index}.xml"))
            for tag in ("cols", "pageMargins", "pageSetup", "mergeCells"):
                assert ET.tostring(new.find(f"s:{tag}", ns)) == ET.tostring(old.find(f"s:{tag}", ns))
        all_xml = b"".join(actual.read(p).lower() for p in actual.namelist() if p.endswith(".xml"))
        assert not any(value in all_xml for value in (b"sanjaybhai", b"parthbhai", b"9924440943", b"krinaben", b"parshantbhai"))


@pytest.mark.parametrize("snapshot,expected", [("9876543210", "9876543210"), ("", ""), (None, "1234567890")])
def test_excel_uses_saved_phone_snapshot_including_explicit_empty(snapshot, expected):
    book = load_workbook(BytesIO(_book({"items": [_line()], "phone_snapshot": snapshot})), data_only=True)
    assert book["Sheet1 (3)"]["D9"].value == expected


def test_prices_use_mrp_and_full_precision_discount_cascade():
    doc = {"items": [_line(), _line(id="l2", room="Guest", qty=3, unit_price=150, mrp=250)],
        "project_discount_pct": 10, "room_discounts": {"Master": {"type": "amount", "value": 123.45}}}
    data = _book(doc)
    values = load_workbook(BytesIO(data), data_only=True)
    formulas = load_workbook(BytesIO(data), data_only=False)
    assert values["BATHROOM 1"]["E3"].value == 1000
    assert values["BATHROOM 1"]["G3"].value == 2000
    assert values["BATHROOM 1"]["K3"].value == pytest.approx(1476.55)
    assert values["BATHROOM 2"]["K3"].value == 405
    assert values["Sheet1 (3)"]["K22"].value == 1881.55
    assert formulas["BATHROOM 1"]["G3"].value == "=E3*F3"
    assert formulas["BATHROOM 1"]["I3"].value == "=E3*H3/100"
    assert formulas["BATHROOM 1"]["J3"].value == "=E3-I3"
    assert formulas["BATHROOM 1"]["K3"].value == "=F3*J3"
    assert formulas["BATHROOM 1"]["F24"].value == "=SUM(F3:F18)"
    assert values["BATHROOM 1"]["B3"].value == "000123"


def test_expands_room_rows_and_summary_without_truncation():
    items = [_line(id=f"line{i}", qty=1) for i in range(23)]
    items += [_line(id=f"room{i}", room=f"Room {i}", qty=1) for i in range(9)]
    data = _book({"items": items})
    formulas = load_workbook(BytesIO(data))
    values = load_workbook(BytesIO(data), data_only=True)
    assert len(formulas.sheetnames) == 11
    assert formulas["BATHROOM 1"]["G26"].value == "=SUM(G3:G25)"
    assert values["BATHROOM 1"]["F26"].value == 23
    assert formulas["Sheet1 (3)"]["K13"].value == "=+'BATHROOM 1'!G26"
    assert values["Sheet1 (3)"]["K24"].value == 25600
    assert "A26:E26" in formulas["BATHROOM 1"].merged_cells
    assert "A23:J23" in formulas["Sheet1 (3)"].merged_cells


def test_user_text_is_literal_and_source_input_is_not_mutated():
    doc = {"customer_name": '=HYPERLINK("https://example.com")', "items": [_line(name="=1+2", sku="+0123")], "created_at": "2026-10-25T23:00:00Z"}
    before = deepcopy(doc)
    book = load_workbook(BytesIO(_book(doc)))
    assert book["Sheet1 (3)"]["D7"].data_type == "s"
    assert book["BATHROOM 1"]["C3"].value == "=1+2"
    assert book["BATHROOM 1"]["C3"].data_type == "s"
    assert book["BATHROOM 1"]["B3"].value == "+0123"
    assert book["Sheet1 (3)"]["D8"].value.day == 26
    assert doc == before


def test_final_additional_costs_reconcile_to_payable():
    data = _book({"items": [_line()], "_po_final": {"total": 1600, "payable_total": 1675,
        "additional_costs": [{"name": "Delivery", "amount": 50}]}})
    book = load_workbook(BytesIO(data), data_only=True)
    assert book["BATHROOM 2"]["A1"].value == "Additional costs"
    assert book["BATHROOM 2"]["C3"].value == "Delivery"
    assert book["BATHROOM 2"]["K24"].value == 75
    assert book["Sheet1 (3)"]["K22"].value == 1675


def test_product_images_are_current_and_embedded(monkeypatch):
    from PIL import Image
    image = BytesIO()
    Image.new("RGB", (40, 20), "red").save(image, format="PNG")
    monkeypatch.setattr(excel, "prefetch_urls", lambda urls, fetch, **_k: [fetch(url) for url in urls])
    monkeypatch.setattr(excel, "remote_image_bytes", lambda _url: image.getvalue())
    data = _book({"items": [_line(image="https://allowed.example/product.png")]})
    with ZipFile(BytesIO(data)) as archive:
        assert "xl/media/quotation_2_3.png" in archive.namelist()
    assert len(load_workbook(BytesIO(data))["BATHROOM 1"]._images) == 1


@pytest.mark.parametrize("doc", [{"doc_type": "tiles_selection"}, {"items": [_line(unit_price=float("inf"))]}])
def test_invalid_document_is_rejected(doc):
    with pytest.raises(ValueError):
        _book(doc)


def test_failed_prefetch_is_not_retried_serially(monkeypatch):
    calls = []
    monkeypatch.setattr(excel, "prefetch_urls", lambda urls, fetch, **_k: [fetch(url) for url in urls])
    monkeypatch.setattr(excel, "remote_image_bytes", lambda url: calls.append(url))
    _book({"items": [_line(image="https://allowed.example/missing.png")]})
    assert calls == ["https://allowed.example/missing.png"]


def test_corrupt_product_image_does_not_discard_quotation(monkeypatch):
    monkeypatch.setattr(excel, "prefetch_urls", lambda urls, fetch, **_k: [fetch(url) for url in urls])
    monkeypatch.setattr(excel, "remote_image_bytes", lambda url: b"not a valid image")
    book = load_workbook(BytesIO(_book({"items": [_line(image="https://allowed.example/corrupt.png")]})), data_only=True)
    assert book["Sheet1 (3)"]["K22"].value == 1600


def test_zero_mrp_uses_selling_base_so_formulas_still_recalculate():
    book = load_workbook(BytesIO(_book({"items": [_line(mrp=0)]})), data_only=True)
    row = book["BATHROOM 1"]
    assert row["E3"].value == 800
    assert row["E3"].value * (1-row["H3"].value/100) * row["F3"].value == row["K3"].value


@pytest.mark.parametrize("line_total,expected_adjustment", [(.34, "-0.01"), (.33, "+0.01")])
def test_final_aggregate_rounding_reconciles_both_directions(line_total, expected_adjustment):
    document = {"items": [_line(id=f"l{i}", qty=1, unit_price=line_total, mrp=1, room=f"Room {i}") for i in range(2)],
        "_po_final": {"total": .67, "payable_total": .67}}
    content = _book(document)
    formulas = load_workbook(BytesIO(content))
    values = load_workbook(BytesIO(content), data_only=True)
    assert values["Sheet1 (3)"]["K22"].value == .67
    assert formulas["Sheet1 (3)"]["K22"].value == "=ROUND(SUM(P13:P20),2)" + expected_adjustment


def test_final_large_total_discrepancy_is_rejected():
    with pytest.raises(ValueError, match="do not reconcile"):
        _book({"items": [_line()], "_po_final": {"total": 2000, "payable_total": 2000}})


def test_xml_invalid_unicode_cannot_corrupt_workbook():
    book = load_workbook(BytesIO(_book({"customer_name": "Client\ud800", "items": [_line(room="Room\uffff")]})))
    assert book["Sheet1 (3)"]["D7"].value == "Client"
    assert book["BATHROOM 1"]["A1"].value == "Room"


def test_cover_logos_are_portable_uncropped_and_keep_their_aspect_ratios():
    from PIL import Image
    data = _book({'items': [_line()]})
    with ZipFile(BytesIO(data)) as archive:
        drawing = ET.fromstring(archive.read('xl/drawings/drawing1.xml'))
        relationships = ET.fromstring(archive.read('xl/drawings/_rels/drawing1.xml.rels'))
        targets = {r.get('Id'): r.get('Target') for r in relationships}
        assert len(drawing) == 11
        for anchor in drawing:
            assert anchor.tag == f'{{{excel.X}}}oneCellAnchor'
            assert anchor.find(f'.//{{{excel.A}}}srcRect') is None
            origin = anchor.find(f'{{{excel.X}}}from')
            assert int(origin.find(f'{{{excel.X}}}colOff').text) < 64 * 9525
            blip = anchor.find(f'.//{{{excel.A}}}blip')
            target = 'xl/' + targets[blip.get(f'{{{excel.R}}}embed')].removeprefix('../')
            with Image.open(BytesIO(archive.read(target))) as image:
                extent = anchor.find(f'{{{excel.X}}}ext')
                assert int(extent.get('cx')) / int(extent.get('cy')) == pytest.approx(image.width / image.height, rel=.00001)
        assert len(load_workbook(BytesIO(data)).worksheets[0]._images) == 11


def test_excel_currency_uses_indian_locale_and_grouping_without_changing_values():
    book = load_workbook(BytesIO(_book({'items': [_line(unit_price=12345678.9, mrp=12345678.9, qty=1)]})), data_only=True)
    for sheet, cell in [('Sheet1 (3)', 'K22'), ('BATHROOM 1', 'E3'), ('BATHROOM 1', 'K3')]:
        assert book[sheet][cell].value == 12345678.9
        assert '[$₹-en-IN]' in book[sheet][cell].number_format
        assert '#,##,##0' in book[sheet][cell].number_format
