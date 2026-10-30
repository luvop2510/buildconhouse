"""Visual exports preserve scope, currency, limits, source rows and user text."""
from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from pypdf import PdfReader

from services.analytics.unified_exports import export_bytes, report_sections


def test_profile_record_only_export_never_silently_truncates():
    with pytest.raises(HTTPException) as failure:
        report_sections({"tables": {"orders": {"items": [{"id": "one"}], "total": 2}}}, view="profile", scope="records")
    assert failure.value.status_code == 422


def test_profile_full_export_contains_selected_tables_and_visual_trend():
    report = {"identity": {"id": "selected-brand", "name": "Selected"}, "kpis": {"sales": 123456.78},
              "trend": [{"id": "day-one", "label": "01 Oct", "sales": 123456.78, "sales_rnb": 0}],
              "tables": {"orders": {"items": [{"id": "source-order", "currency": "INR", "sales": 123456.78}], "total": 1}}}
    sections = report_sections(report, view="profile")
    assert "Trend" in [name for name, _ in sections]
    data = export_bytes(report, sections, "pdf", {"Selected entity ID": "selected-brand"})
    text = "\n".join(page.extract_text() for page in PdfReader(BytesIO(data)).pages)
    assert "INR 1,23,456.78" in text
    assert "selected-brand" in text and "source-order" in text and "01 Oct" in text
    assert "INR / sales over time" in text and "Bar =" not in text
    assert len(PdfReader(BytesIO(data)).pages) == 1


def test_workbook_preserves_numeric_money_and_percentage_units():
    sections = [("Summary", [{"metric": "sales rnb", "value": 1234.56}, {"metric": "change pct", "value": 12.5}]),
                ("Lines", [{"id": "line", "currency": "INR", "sales": 123456.78, "discount_pct": 10, "name": "=not-a-formula"}])]
    book = load_workbook(BytesIO(export_bytes({}, sections, "xlsx", {})))
    assert book["Summary"]["B2"].value == 1234.56
    assert '"RNB "' in book["Summary"]["B2"].number_format
    assert '%' in book["Summary"]["B3"].number_format
    assert book["Lines"]["E2"].data_type == "s"
    assert book["Lines"]["E2"].value.startswith("'=")
    assert '"INR "' in book["Lines"]["C2"].number_format


def test_long_untrusted_values_are_printable_without_markup_or_lost_rows():
    data = export_bytes({"title": "Customer <script> & history"}, [("History", [{"id": "long-text-row", "notes": "<b>literal</b> " * 600}])], "pdf", {})
    text = "\n".join(page.extract_text() for page in PdfReader(BytesIO(data)).pages)
    assert "Customer <script> & history" in text and "long-text-row" in text
    assert "<b>literal</b>" in text


def test_large_workbook_formats_without_scanning_all_existing_cells_per_cell(monkeypatch):
    """Formatting work must stay linear as the permitted export grows."""
    from openpyxl.worksheet.worksheet import Worksheet
    original = Worksheet.max_row.fget
    calls = 0
    def counted_max_row(sheet):
        nonlocal calls
        calls += 1
        return original(sheet)
    monkeypatch.setattr(Worksheet, "max_row", property(counted_max_row))
    rows = [{"id": f"order-{index}", "name": "Recorded product", "sales": index + 100.25, "currency": "INR", "orders": 1, "customers": 1,
             "quantity": 1.2345, "unit": "Pieces", "date": "2026-10-28", "company_id": "brand1", "customer_id": "customer1", "discount_pct": 10} for index in range(3000)]
    blob = export_bytes({"title": "Large profile"}, [("Products", rows)], "xlsx", {"Date basis": "confirmation"})
    # Serialization may inspect dimensions; formatting may not scan the entire
    # worksheet for every one of the 36,000 cells.
    assert calls < 50
    book = load_workbook(BytesIO(blob))
    sheet = book["Products"]
    assert sheet.max_row == 3001 and sheet.max_column == 12
    assert sheet["A2"].value == "order-0" and sheet["A3001"].value == "order-2999"
    assert sheet["C3001"].value == 3099.25 and '"INR "' in sheet["C3001"].number_format
    assert sheet["L3001"].value == 10 and sheet["L3001"].number_format == '0.00"%"'
    assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:L3001"


def test_pdf_is_portrait_a4_and_keeps_every_continuation_row():
    from reportlab.lib.pagesizes import A4
    rows = [{"id": f"order-{index}", "name": "Recorded customer", "sales": 1000 + index, "currency": "INR", "orders": 1} for index in range(55)]
    reader = PdfReader(BytesIO(export_bytes({"title": "Selected sales"}, [("Orders", rows)], "pdf", {"Period": "October 2026"})))
    assert len(reader.pages) > 1
    for page in reader.pages:
        assert float(page.mediabox.width) == pytest.approx(A4[0], abs=.01)
        assert float(page.mediabox.height) == pytest.approx(A4[1], abs=.01)
        assert "Amounts show their currency" in page.extract_text()
    text = "\n".join(page.extract_text() for page in reader.pages)
    assert all(f"order-{index}" in text for index in range(55))
    assert "Largest recorded values" not in text


def test_records_only_pdf_does_not_add_unselected_dashboard_totals():
    report = {"kpis": {"sales": 987654321}, "trend": [{"label": "Hidden month", "sales": 987654321}],
              "tables": {"orders": {"items": [{"id": "selected-record", "sales": None}], "total": 1}}}
    sections = report_sections(report, view="profile", scope="records")
    text = "\n".join(page.extract_text() for page in PdfReader(BytesIO(export_bytes(report, sections, "pdf", {}))).pages)
    assert "selected-record" in text and "Unavailable" in text
    assert "98,76,54,321" not in text and "Hidden month" not in text


def test_pdf_keeps_exact_currency_and_scope_without_redundant_dashboard_content():
    sections = [("Orders", [{"id": "catalog-order", "sales": 123.45, "currency": "INR"},
                             {"id": "furniture-order", "sales": 678.9, "currency": "RNB"}])]
    report = {"title": "Selected sales", "insights": ["Unselected dashboard narrative"]}
    metadata = {"Period": "October 2026", "Date basis": "payment", "Generated at": "internal-generation-time",
                "Export scope": "records", "Filters": {"floor_id": "all", "currency": "INR", "customer_id": "selected-customer", "sort": "sales_desc", "granularity": "month"}}
    reader = PdfReader(BytesIO(export_bytes(report, sections, "pdf", metadata)))
    text = "\n".join(page.extract_text() for page in reader.pages)
    assert "INR 123.45" in text and "RNB 678.9" in text
    assert "Period: October 2026" in text and "Date basis: payment" in text and "selected-customer" in text
    assert "Department: All departments" in text and "Floor id" not in text
    assert "Unselected dashboard narrative" not in text and "internal-generation-time" not in text
    assert "sales_desc" not in text and "granularity" not in text and "Authorized report" not in text
    assert text.count("catalog-order") == 1 and text.count("furniture-order") == 1
    assert len(reader.pages) == 1


def test_current_profile_tab_and_record_directory_export_only_selected_data():
    report = {"kpis": {"sales": 987654321}, "identity": {"id": "brand", "name": "Brand"},
              "tables": {"orders": {"items": [{"id": "selected-order"}], "total": 1},
                         "payments": {"items": [{"id": "hidden-payment"}], "total": 1}}}
    assert report_sections(report, view="profile", scope="current", tab="orders") == [("Orders", [{"id": "selected-order"}])]
    assert report_sections({"kpis": {"sales": 987654321}}, [{"id": "selected-order"}], view="records") == [("Records", [{"id": "selected-order"}])]


def test_pdf_retains_missing_values_and_source_limitations_as_data_notes():
    report = {"coverage": {"notes": ["Some historical line totals are unavailable."]}}
    sections = [("Trend", [{"label": "One", "sales": 100}, {"label": "Unknown", "sales": None}, {"label": "Three", "sales": 200}])]
    text = "\n".join(page.extract_text() for page in PdfReader(BytesIO(export_bytes(report, sections, "pdf", {}))).pages)
    assert "Unknown" in text and "Unavailable" in text and "Some historical line totals are unavailable." in text
    assert "INR 0" not in text


def test_visual_overview_is_compact_and_keeps_all_selected_orders():
    from pypdf.generic import ContentStream
    report = {"title": "October sales review", "filters": {"floor_id": "all"},
        "kpis": {"sales": 300000.5, "orders": 3, "customers": 3, "average_order_value": 150000.25, "collected": 220000, "outstanding": 80000.5, "products": 3, "change_pct": None, "sales_rnb": 75000.5, "collected_rnb": 62500, "outstanding_rnb": 12500.5},
        "trend": [{"id": f"date-{i}", "label": f"{day} Oct", "sales": value, "orders": 1, "previous_sales": None} for i, (day, value) in enumerate([(10, 240000), (20, 60000.5), (23, 0)])],
        "floors": [{"name": name, "floor_id": floor, "currency": currency, "sales": value, "orders": 1, "customers": 1, "contribution_pct": 100} for name, floor, currency, value in [("Sanitary", "first-floor", "INR", 240000), ("Tiles", "ground-floor", "INR", 60000.5), ("Furniture", "third-floor", "RNB", 75000.5)]],
        "customer_segments": [{"name": "New customers", "customers": 3}, {"name": "Repeat customers", "customers": 0}],
        "coverage": {"notes": ["Previous-period comparison unavailable: no reliable historical baseline."]}}
    for key in ("companies", "products", "architects", "customers"):
        report[key] = [{"id": f"{key}-source", "name": f"Recorded {key}", "sales": 300000.5, "orders": 3, "customers": 3}]
    orders = [{"id": f"selected-source-{i}", "number": f"SO-{i}", "customer_name": "Recorded customer", "floor_name": "Sanitary", "status": "confirmed", "sales": 100000, "collected": 50000, "outstanding": 50000, "date": "2026-10-10T06:00:00Z", "currency": "INR"} for i in range(3)]
    reader = PdfReader(BytesIO(export_bytes(report, report_sections(report, orders), "pdf", {"Period": "October 2026", "Date basis": "confirmation"})))
    assert len(reader.pages) <= 2
    text = "\n".join(page.extract_text() for page in reader.pages)
    assert all(order["id"] in text for order in orders)
    assert "INR 3,00,000.50" in text and "RNB 75,000.50" in text
    assert "INR / sales over time" in text and "10 Oct 2026" in text
    assert "no reliable historical baseline" in text
    operations = [operation for page in reader.pages for _, operation in ContentStream(page.get_contents(), reader).operations]
    assert operations.count(b"re") >= 10


def test_furniture_customer_export_keeps_its_actual_rnb_amount():
    report = {"filters": {"floor_id": "all"}, "customers": [{"id": "furniture-customer", "name": "Furniture customer", "floor_ids": ["third-floor"], "sales": 0, "sales_rnb": 75000.5, "outstanding": 12500.5}]}
    sections = report_sections(report)
    row = dict(sections)["Customers"][0]
    assert row["sales"] == 75000.5 and row["currency"] == "RNB"
    reader = PdfReader(BytesIO(export_bytes(report, sections, "pdf", {})))
    text = "\n".join(page.extract_text() for page in reader.pages)
    assert "RNB 75,000.50" in text and "RNB 0.00" not in text


def test_furniture_compact_profile_keeps_primary_sales_and_current_balance():
    report = {"filters": {"floor_id": "third-floor"}, "customers": [{"id": "furniture-source", "name": "Furniture customer", "floor_ids": ["third-floor"], "sales": 0, "sales_rnb": 75000.5, "outstanding": 0, "outstanding_rnb": 12500.5, "customers": 1, "orders": 3}]}
    reader = PdfReader(BytesIO(export_bytes(report, report_sections(report), "pdf", {})))
    text = "\n".join(page.extract_text() for page in reader.pages)
    assert "RNB 75,000.50" in text and "RNB 12,500.50" in text
    assert "Unavailable" not in text.replace("Unavailable is not zero.", "") and "INR" not in text


def test_payment_report_chart_labels_cash_as_collections():
    report = {"filters": {"floor_id": "first-floor", "date_basis": "payment"}, "trend": [{"label": "10 Oct", "sales": 1500, "previous_sales": None, "orders": 1}]}
    reader = PdfReader(BytesIO(export_bytes(report, report_sections(report), "pdf", {"Date basis": "payment"})))
    text = "\n".join(page.extract_text() for page in reader.pages)
    assert "INR / collections over time" in text and "INR / sales over time" not in text


def test_pdf_collection_arcs_use_only_complete_selected_currency_balances():
    report = {"filters": {"floor_id": "all"}, "kpis": {"sales": 400, "collected": 300, "outstanding": 100,
              "sales_rnb": 100, "collected_rnb": 25, "outstanding_rnb": 75}, "floors": [{"currency": "RNB"}]}
    reader = PdfReader(BytesIO(export_bytes(report, report_sections(report), "pdf", {})))
    text = "\n".join(page.extract_text() for page in reader.pages)
    assert "INR collected share" in text and "75.0%" in text
    assert "RNB collected share" in text and "25.0%" in text
    assert "Receipts / (receipts + outstanding)" in text and "target" not in text.lower()
    assert len(reader.pages) == 1
    scoped = {**report, "tables": {"orders": {"items": [{"id": "only-order", "sales": 400}], "total": 1}}}
    text = "\n".join(page.extract_text() for page in PdfReader(BytesIO(export_bytes(scoped, report_sections(scoped, view="profile", scope="records"), "pdf", {}))).pages)
    assert "only-order" in text and "collected share" not in text


@pytest.mark.parametrize("collected,outstanding", [(None, 100), (100, None), (0, 0), (-10, 100)])
def test_pdf_does_not_invent_collection_progress_for_incomplete_or_incomparable_balances(collected, outstanding):
    report = {"filters": {"floor_id": "first-floor"}, "kpis": {"collected": collected, "outstanding": outstanding}}
    text = "\n".join(page.extract_text() for page in PdfReader(BytesIO(export_bytes(report, report_sections(report), "pdf", {}))).pages)
    assert "collected share" not in text


def test_pdf_omits_pagination_metadata_but_keeps_selected_scope():
    report = {"filters": {"floor_id": "first-floor", "page": "3", "page_size": "30", "customer_id": "selected-client"}}
    text = "\n".join(page.extract_text() for page in PdfReader(BytesIO(export_bytes(report, [("Orders", [{"id": "selected-order"}])], "pdf", {}))).pages)
    assert "selected-client" in text and "selected-order" in text
    assert "Page size" not in text and "Page: 3" not in text


def test_record_pdf_compacts_transport_columns_without_losing_source_attribution():
    row = {"id": "record-source", "number": "SO-2026", "customer_id": "internal-customer-id", "customer_name": "Client name",
           "floor_id": "first-floor", "currency": "INR", "date": "2026-10-10T18:45:00Z", "confirmation_date": "2026-10-08T06:00:00Z",
           "status": "confirmed", "sales": 100000.25, "collected": 75000, "outstanding": 25000.25,
           "architect_id": "internal-architect-id", "architect_name": "Architect name", "source": "sanitary"}
    reader = PdfReader(BytesIO(export_bytes({}, [("Records", [row])], "pdf", {})))
    text = "\n".join(page.extract_text() for page in reader.pages)
    for value in ("record-source", "SO-2026", "Client name", "Architect name", "11 Oct 2026", "Confirmed: 08 Oct 2026", "sanitary", "confirmed", "INR 1,00,000.25", "INR 75,000.00", "INR 25,000.25"):
        assert value in text
    assert "internal-customer-id" not in text and "internal-architect-id" not in text
    assert "detail columns" not in text and len(reader.pages) == 1


def test_pdf_trend_identifies_order_counts_across_selected_currencies():
    report = {"filters": {"floor_id": "all"}, "trend": [{"label": "23 Oct", "sales": 0, "sales_rnb": 75000.5, "orders": 1, "orders_inr": 0, "orders_rnb": 1}]}
    sections = report_sections(report)
    text = " ".join(" ".join(page.extract_text().split()) for page in PdfReader(BytesIO(export_bytes(report, sections, "pdf", {}))).pages)
    assert "Orders (all selected currencies)" in text and "INR 0.00" in text
    assert dict(sections)["Trend"][0]["orders"] == 1
    book = load_workbook(BytesIO(export_bytes(report, sections, "xlsx", {})))
    sheet = book["Trend"]
    headers = [entry.value for entry in sheet[1]]
    assert sheet.cell(2, headers.index("orders") + 1).value == 1
    assert "Orders (all selected currencies)" not in headers


def test_product_pdf_keeps_crore_money_on_one_rendered_line_without_reducing_type():
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.platypus import Paragraph
    from services.analytics.unified_exports import PDF_WIDTH, _pdf_bands, _pdf_rows, display
    from pdf_generator import RUPEE_FONT
    rows = [{"id": f"product-{i}", "name": "Recorded product with a descriptive catalog name", "currency": "INR",
             "sales": amount, "orders": 6, "customers": 4, "outstanding": None, "quantity": 160, "unit": "Pieces", "company_name": "Recorded company"}
            for i, amount in enumerate((79999920.0, 19167390.0, 999999999.99))]
    normalized = _pdf_rows("Products", rows)
    face = RUPEE_FONT or "Helvetica"
    bands = _pdf_bands(normalized, "name", [key for key in normalized[0] if key not in {"name", "currency"}], PDF_WIDTH, face)
    for band, widths in bands:
        assert sum(widths) == pytest.approx(PDF_WIDTH)
        assert widths[0] >= 140
        if "sales" in band:
            content_width = widths[band.index("sales")] - 14
            for row in rows:
                value = display(row["sales"], "sales", "INR")
                assert content_width >= stringWidth(value, face, 8)
                paragraph = Paragraph(value, ParagraphStyle("MeasuredMoney", fontName=face, fontSize=8, leading=10, wordWrap="CJK"))
                paragraph.wrap(content_width, 1000)
                assert len(paragraph.blPara.lines) == 1
    reader = PdfReader(BytesIO(export_bytes({}, [("Products", rows)], "pdf", {})))
    runs = []
    for page in reader.pages:
        page.extract_text(visitor_text=lambda text, cm, tm, font, size: runs.append((text.strip(), size)))
    for row in rows:
        value = display(row["sales"], "sales", "INR")
        assert any(text == value and size == 8 for text, size in runs)


@pytest.mark.parametrize("section", ["Source Details", "Orders", "Records"])
def test_transaction_pdf_keeps_all_crore_amounts_on_single_rendered_lines(section):
    from services.analytics.unified_exports import display
    row = {"id": "source-record", "number": "SO-2026", "customer_name": "Recorded customer", "floor_id": "first-floor", "currency": "INR",
           "sales": 999999999.99, "collected": 199999999.99, "outstanding": 800000000.0}
    reader = PdfReader(BytesIO(export_bytes({}, [(section, [row])], "pdf", {})))
    runs = []
    for page in reader.pages:
        page.extract_text(visitor_text=lambda text, cm, tm, font, size: runs.append((text.strip(), size)))
    for key in ("sales", "collected", "outstanding"):
        assert any(text == display(row[key], key, "INR") and size == 8 for text, size in runs)
