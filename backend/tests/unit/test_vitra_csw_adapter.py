from io import BytesIO

from openpyxl import Workbook

from catalog_pipeline.adapters.vitra import VitraAdapter
from catalog_pipeline.base import MISSING


def _workbook_bytes() -> bytes:
    workbook = Workbook()
    lookup = workbook.active
    lookup.title = "Sheet1"
    lookup.append(["DESIGN", "DETAIL", "WHITE 003", None, None])
    lookup.append([None, None, "Code", "IMAGE", "MRP"])
    lookup.append(["LOOKUP ONLY", "WC", "LOOKUPB403H0001", None, 999])

    csw = workbook.create_sheet("CSW")
    csw.append(["DESIGN", "DETAIL", "WHITE 003", None, None])
    csw.append([None, None, "Code", "IMAGE", "MRP"])
    # openpyxl does not write a calculated cache.  This represents a VLOOKUP
    # workbook that has not been recalculated and must remain review-only.
    csw.append(["MEMORIA", "RIM-EX WC", "5885B403H0101", None, "=VLOOKUP(C3,Sheet1!A:E,5,FALSE)"])

    extra = workbook.create_sheet("Sheet2")
    extra.append(["DESIGN", "DETAIL", "WHITE 003", None, None])
    extra.append([None, None, "Code", "IMAGE", "MRP"])
    extra.append(["NOT A PRODUCT", "WC", "EXTRAB403H0001", None, 111])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def test_csw_reimport_adapter_excludes_lookup_and_other_sheets():
    rows, report = VitraAdapter(only_sheets={"CSW"}).extract(_workbook_bytes(), "vitra.xlsx")

    assert [row.sku for row in rows] == ["5885B403H0101"]
    assert rows[0].family_key.startswith("vitra:csw:")
    assert rows[0].mrp == MISSING
    assert any("no cached MRP value" in issue for issue in rows[0].issues)
    assert any("CSW!E3" in warning for warning in report.warnings)


def test_default_vitra_adapter_is_not_sheet_restricted():
    rows, _ = VitraAdapter().extract(_workbook_bytes(), "vitra.xlsx")

    assert {row.sku for row in rows} == {"LOOKUPB403H0001", "5885B403H0101", "EXTRAB403H0001"}


def test_csw_wc_bidets_receive_a_source_derived_configuration_group():
    workbook = Workbook()
    ws = workbook.active
    ws.title = "CSW"
    ws.append(["DESIGN", "DETAIL", "WHITE 003", None, None])
    ws.append([None, None, "Code", "IMAGE", "MRP"])
    ws.append(["MEMORIA", "RIM-EX WC", "5885B403H0101", None, 100])
    ws.append([None, "RIM-EX WC WITH BIDET", "5885B403H0559", None, 110])
    stream = BytesIO()
    workbook.save(stream)

    rows, _ = VitraAdapter(only_sheets={"CSW"}).extract(stream.getvalue(), "vitra.xlsx")

    assert rows[0].specs["configuration_label"] == "Without bidet"
    assert rows[1].specs["configuration_label"] == "With bidet"
    assert rows[0].specs["configuration_group"] == rows[1].specs["configuration_group"]
