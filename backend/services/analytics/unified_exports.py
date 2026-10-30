"""Server-derived PDF and Excel exports of exactly the authorized report view."""
from __future__ import annotations

from io import BytesIO
from datetime import datetime, timezone
from zoneinfo import ZoneInfo
from xml.sax.saxutils import escape

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

from services.analytics.unified import FLOORS
from services.analytics.unified_pdf_graphics import numeric, trend, ranking, cell_bar, collection_arc
from services.export import spreadsheet_cell
from services.money import format_indian_money
from pdf_generator import RUPEE_FONT

MAX_EXPORT_ROWS = 10_000
ACCENT = "0066CC"
MONEY_FIELDS = {"sales", "sales_rnb", "average_order_value", "average_order_value_rnb", "collected", "outstanding", "collected_rnb", "outstanding_rnb", "unit_price", "line_total", "amount", "subtotal", "total", "discount", "tax", "taxes"}

PDF_MONEY_FIELDS = MONEY_FIELDS | {"previous_sales", "previous_sales_rnb"}


def display(value, key="", currency="INR"):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if key in PDF_MONEY_FIELDS:
            return f"{'RNB' if key.endswith('_rnb') or currency == 'RNB' else 'INR'} {format_indian_money(value)}"
        if key.endswith("_pct") or key.endswith("_pct_rnb"):
            return f"{value:.1f}%"
    return str(cell(value))


# PDF colors follow the app's semantic surface and content roles.
PDF_INK = colors.HexColor("#1D1D1F")
PDF_SECONDARY = colors.HexColor("#545458")
PDF_LINE = colors.HexColor("#E2E2E7")
PDF_WASH = colors.HexColor("#F5F5F7")
PDF_WIDTH = A4[0] - 72


def cell(value):
    if value is None:
        return "Unavailable"
    if isinstance(value, list):
        return "; ".join(str(item) for item in value)
    if isinstance(value, dict):
        return "; ".join(f"{key}: {item}" for key, item in value.items())
    return value


def report_sections(report, records=None, *, view="dashboard", scope="full", tab=None):
    result = []
    furniture = report.get("filters", {}).get("floor_id") == "third-floor"
    inr_fields = {"sales", "average_order_value", "collected", "outstanding", "orders_inr", "previous_sales", "change_pct"}
    def visible(key, aggregate=False):
        if key.endswith("_rnb"):
            return furniture
        return not (furniture and aggregate and key in inr_fields)

    summary = report.get("kpis") or report.get("summary")
    overview = view not in {"companies", "products", "architects", "customers", "records"} and (scope == "full" or scope == "current" and tab == "overview")
    if overview and summary:
        result.append(("Summary", [{"metric": key.replace("_", " "), "value": value} for key, value in summary.items() if visible(key, "sales_rnb" in summary)]))
        if not furniture and any(floor.get("currency") == "RNB" for floor in report.get("floors", [])):
            result.append(("Furniture ledger", [{"metric": key.replace("_", " "), "value": value} for key, value in summary.items() if key.endswith("_rnb")]))
    if "identity" in report and overview:
        identity = report["identity"]
        result.append(("Identity", [{"field": "Stable ID", "value": identity["id"]}, {"field": "Name", "value": identity["name"]}, *[{"field": a["label"], "value": a["value"]} for a in identity.get("attributes", [])]]))
    if report.get("tables") and scope == "current" and tab == "overview":
        for name in ("trend", "floors", "companies", "architects", "products", "customers"):
            result.append((name.title(), report.get(name, [])))
    elif report.get("tables"):
        if scope == "records":
            table = report["tables"]["orders"]
            if table["total"] > len(table["items"]):
                raise HTTPException(422, "This export exceeds its row limit; narrow the filters")
            result.append(("Orders", table["items"]))
        else:
            if scope == "full":
                for name in ("trend", "floors", "companies", "architects", "products", "customers"):
                    result.append((name.title(), report.get(name, [])))
                for name, values in report.get("attribute_breakdowns", {}).items():
                    result.append((name.title() + " Analysis", values))
            for name, table in report["tables"].items():
                if scope == "current" and tab and name != tab:
                    continue
                if table["total"] > len(table["items"]):
                    raise HTTPException(422, "This export exceeds its row limit; narrow the filters")
                result.append((name.replace("_", " ").title(), table["items"]))
    elif "lines" in report:
        for name in ("lines", "payments", "timeline", "documents"):
            result.append((name.title(), report.get(name, [])))
    elif view in {"companies", "products", "architects", "customers", "records"}:
        result.append((view.title(), records or []))
    else:
        for name in ("trend", "floors", "companies", "architects", "products", "customers", "customer_segments"):
            result.append((name.replace("_", " ").title(), report.get(name, [])))
        for name, values in report.get("attribute_breakdowns", {}).items():
            result.append((name.title() + " Analysis", values))
        if records is not None:
            result.append(("Source Details", records))
    # Aggregate columns follow the visible currency; actual source records keep
    # their explicit currency, including Furniture rows in company-wide reports.
    def selected_currency(row):
        if row.get("floor_ids") == ["third-floor"] and "sales_rnb" in row:
            row = {**row, "sales": row["sales_rnb"], "currency": "RNB"}
        return {key: value for key, value in row.items() if visible(key, "sales_rnb" in row)}
    result = [(name, [selected_currency(row) for row in rows]) for name, rows in result]
    if sum(len(rows) for _, rows in result) > MAX_EXPORT_ROWS:
        raise HTTPException(422, "Export exceeds 10,000 rows; narrow the period, floor or entity filters")
    return result


def export_bytes(report, sections, fmt, metadata):
    if fmt == "xlsx":
        return workbook(report, sections, metadata)
    return pdf(report, sections, metadata)


def workbook(report, sections, metadata):
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    for key, value in metadata.items():
        ws.append([spreadsheet_cell(key), spreadsheet_cell(str(cell(value)))])
    for insight in report.get("insights", []):
        ws.append(["Data brief", spreadsheet_cell(insight)])
    for note in report.get("coverage", {}).get("notes", report.get("notes", [])):
        ws.append(["Coverage", spreadsheet_cell(note)])
    for name, rows in sections:
        sheet = wb.create_sheet(name[:31])
        keys = list(dict.fromkeys(key for row in rows for key in row))
        sheet.append(keys or ["No matching records"])
        for row_number, row in enumerate(rows, 2):
            sheet.append([spreadsheet_cell(cell(row.get(key))) for key in keys])
            for index, key in enumerate(keys, 1):
                value = row.get(key)
                effective_key = row.get("metric", "").replace(" ", "_") if key == "value" else key
                target = sheet.cell(row_number, index)
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    if effective_key in MONEY_FIELDS:
                        currency = "RNB" if effective_key.endswith("_rnb") or row.get("currency") == "RNB" else "INR"
                        target.number_format = f'"{currency} "#,##,##0.00;[Red]-"{currency} "#,##,##0.00'
                    elif effective_key.endswith("_pct"):
                        target.number_format = '0.00"%"'
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for header in sheet[1]:
            header.font = Font(bold=True, color="FFFFFF")
            header.fill = PatternFill("solid", fgColor=ACCENT)
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(44, max(16, max(len(str(c.value or "")) for c in column[:100]) + 2))
        if name == "Trend" and len(rows) > 1:
            for position, (key, currency) in enumerate((("sales", "INR"), ("sales_rnb", "RNB"))):
                if key not in keys or not any(row.get(key) for row in rows):
                    continue
                chart = BarChart()
                chart.title = report.get("metric_label", "Confirmed sales") + f" ({currency})"
                chart.y_axis.title = currency
                chart.add_data(Reference(sheet, min_col=keys.index(key) + 1, min_row=1, max_row=len(rows) + 1), titles_from_data=True)
                if "label" in keys:
                    chart.set_categories(Reference(sheet, min_col=keys.index("label") + 1, min_row=2, max_row=len(rows) + 1))
                sheet.add_chart(chart, f"S{2 + position * 17}")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 100
    data = BytesIO()
    wb.save(data)
    return data.getvalue()


def _label(key):
    label = key.replace("_rnb", " (RNB)").replace("_", " ").capitalize()
    return "ID" if key == "id" else label.replace("(rnb)", "(RNB)").replace(" id", " ID")


def _pdf_date(value):
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        return (parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)).astimezone(ZoneInfo("Asia/Kolkata")).strftime("%d %b %Y")
    except (TypeError, ValueError):
        return str(value or "Date unavailable")


def _pdf_rows(name, rows):
    """Drop duplicate transport fields, retaining source identity and useful data."""
    if name == "Trend":
        fields = {"label", "sales", "sales_rnb", "previous_sales", "previous_sales_rnb", "orders", "quantity", "unit"}
    elif name == "Floors":
        fields = {"name", "sales", "orders", "customers", "contribution_pct"}
    elif name in {"Companies", "Architects", "Products", "Customers"}:
        fields = {"name", "sales", "sales_rnb", "orders", "customers", "quantity", "unit", "sku", "company_name", "outstanding", "outstanding_rnb"}
    elif name == "Customer Segments":
        fields = {"name", "customers"}
    elif name in {"Source Details", "Orders", "Records"}:
        fields = {"number", "sales", "collected", "outstanding"}
    else:
        return rows
    result = []
    for source in rows:
        row = {key: value for key, value in source.items() if key in fields}
        row["currency"] = source.get("currency", "RNB" if source.get("floor_ids") == ["third-floor"] else "INR")
        if name in {"Companies", "Architects", "Products", "Customers"} and source.get("id"):
            row["name"] = f"{source.get('name', 'Recorded identity')}\n{source['id']}"
        if name in {"Source Details", "Orders", "Records"}:
            stamp = _pdf_date(source.get("date"))
            row["number"] = "\n".join(filter(None, [
                " · ".join(str(value) for value in (source.get("number") or source.get("id"), source.get("customer_name") or source.get("name")) if value),
                " · ".join(str(value) for value in (source.get("floor_name") or FLOORS.get(source.get("floor_id"), source.get("floor_id")), source.get("status"), stamp) if value),
                " · ".join(str(value) for value in (source.get("architect_name"), source.get("source"), source.get("id") if source.get("number") and source.get("number") != source.get("id") else None) if value),
                f"Confirmed: {_pdf_date(source['confirmation_date'])}" if source.get("confirmation_date") and _pdf_date(source['confirmation_date']) != stamp else "",
            ]))
        if "quantity" in row and "unit" in row:
            row["quantity"] = f"{row['quantity']} {row.pop('unit') or ''}" if row['quantity'] is not None else None
        if name == "Products" and row.get("sku"):
            row["name"] = f"{row.get('name', '')} / {row.pop('sku')}"
        if name == "Customers":
            row.pop("customers", None)
        result.append(row)
    return result


def pdf(report, sections, metadata):
    blob = BytesIO()
    styles = getSampleStyleSheet()
    face = RUPEE_FONT or "Helvetica"
    for style in styles.byName.values():
        style.fontName = face
    styles.add(ParagraphStyle("ReportTitle", fontName=face, fontSize=22, leading=27, textColor=PDF_INK, spaceAfter=10))
    styles.add(ParagraphStyle("SectionTitle", fontName=face, fontSize=12, leading=16, textColor=PDF_INK, spaceBefore=10, spaceAfter=6, keepWithNext=True))
    styles.add(ParagraphStyle("SmallReport", fontName=face, fontSize=8, leading=11, textColor=PDF_SECONDARY, spaceAfter=3, wordWrap="CJK"))
    styles.add(ParagraphStyle("DetailCaption", parent=styles["SmallReport"], keepWithNext=True))
    styles.add(ParagraphStyle("CellReport", fontName=face, fontSize=8, leading=10, textColor=PDF_INK, wordWrap="CJK"))
    styles.add(ParagraphStyle("NumberReport", parent=styles["CellReport"], alignment=2))
    small, body = styles["SmallReport"], styles["CellReport"]
    para = lambda value, style=body: Paragraph(escape(str(value)).replace("\n", "<br/>"), style)
    title = report.get("title") or "Sales analytics"
    story = [para(title, styles["ReportTitle"])]
    for key in ("Period", "Date basis", "Date from", "Date to (exclusive)", "Directory"):
        if metadata.get(key) is not None:
            story.append(para(f"{key}: {metadata[key]}", small))
    filters = metadata.get("Filters") or report.get("filters") or {}
    # Presets and presentation settings repeat the dates or affect ordering only.
    active = {key: value for key, value in filters.items() if value is not None and value != "" and key not in {"preset", "date_basis", "date_from", "date_to", "sort", "granularity", "page", "page_size"}}
    if metadata.get("Floor") is not None:
        active.setdefault("floor_id", metadata["Floor"])
    selected_id = metadata.get("Selected entity ID")
    if selected_id and selected_id not in active.values() and selected_id != report.get("identity", {}).get("id"):
        active["selected_entity_id"] = selected_id
    if "floor_id" in active:
        floor = active.pop("floor_id")
        active = {"department": "All departments" if floor == "all" else FLOORS.get(floor, floor), **active}
    if active:
        story.append(para(" / ".join(f"{_label(key)}: {value}" for key, value in active.items()), small))
    story.append(Spacer(1, 8))
    empty_sections = []
    table_style = [("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5)]
    pending = []
    def flush_panels():
        if not pending:
            return
        cells = pending[:]
        if len(cells) == 1:
            cells.append("")
        grid = Table([cells], colWidths=[PDF_WIDTH / 2] * 2, hAlign="LEFT")
        grid.setStyle(TableStyle(table_style + [("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
        story.append(grid)
        pending.clear()
    for name, rows in sections:
        rows = _pdf_rows(name, rows)
        if not rows:
            empty_sections.append(name)
            continue
        paired = name in {"Floors", "Companies", "Architects", "Products", "Customers", "Customer Segments"} and len(rows) <= 8
        if not paired:
            flush_panels()
        section_width = PDF_WIDTH / 2 - 14 if paired else PDF_WIDTH
        block = [para(name, styles["SectionTitle"])]
        if not paired:
            story.extend(block)
        if name in {"Summary", "Furniture ledger"}:
            tiles = []
            metrics = {row.get("metric", "").replace(" ", "_"): row.get("value") for row in rows}
            ledger = "RNB" if name == "Furniture ledger" or filters.get("floor_id") == "third-floor" else "INR"
            suffix = "_rnb" if ledger == "RNB" else ""
            arc = collection_arc(metrics.get(f"collected{suffix}", metrics.get("collected")), metrics.get(f"outstanding{suffix}", metrics.get("outstanding")), ledger, 150)
            for row in rows:
                metric = row.get("metric", "").replace(" ", "_")
                if metric.startswith("change_pct") and row.get("value") is None:
                    continue
                tiles.append([para(_label(metric), small), para(display(row.get("value"), metric), styles["CellReport"])])
            if not tiles:
                story.append(para("Comparable growth unavailable", small))
                continue
            columns = 3 if arc else 4
            card_width = PDF_WIDTH - 174 if arc else PDF_WIDTH
            grid = [tiles[i:i + columns] + [""] * (columns - len(tiles[i:i + columns])) for i in range(0, len(tiles), columns)]
            cards = Table(grid, colWidths=[card_width / columns] * columns, hAlign="LEFT", cornerRadii=[10] * 4)
            cards.setStyle(TableStyle(table_style + [("BACKGROUND", (0, 0), (-1, -1), PDF_WASH), ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
            if arc:
                summary = Table([[arc, cards]], colWidths=[164, PDF_WIDTH - 164], hAlign="LEFT", cornerRadii=[10] * 4)
                summary.setStyle(TableStyle(table_style + [("BACKGROUND", (0, 0), (-1, -1), PDF_WASH), ("BOX", (0, 0), (-1, -1), .5, PDF_LINE), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
                story.append(summary)
            else:
                story.append(cards)
            continue
        keys = [key for key in dict.fromkeys(key for row in rows for key in row) if key != "currency"]
        if not keys:
            continue
        identity = next((key for key in ("name", "number", "label", "metric", "field", "id", "review_id") if key in keys), keys[0])
        if paired:
            metric = "sales" if any("sales" in row for row in rows) else "sales_rnb" if any("sales_rnb" in row for row in rows) else "customers"
            values = []
            for row in rows:
                currency = row.get("currency", "INR")
                extras = " · ".join(f"{_label(key)}: {display(value, key, currency)}" for key, value in row.items() if key not in {identity, metric, "currency"})
                label = [para(row.get(identity, "Recorded row")), para(extras, small)]
                raw = row.get(metric)
                peers = [item.get(metric) for item in rows if numeric(item.get(metric)) and item.get("currency", "INR") == currency]
                value = [para(display(raw, metric, currency), styles["NumberReport"])]
                if numeric(raw) and peers:
                    value.append(cell_bar(raw, min(0, *peers), max(0, *peers), section_width * .4 - 10))
                values.append([label, value])
            compact_table = Table(values, colWidths=[section_width * .6, section_width * .4], hAlign="LEFT", splitInRow=1)
            compact_table.setStyle(TableStyle(table_style + [("LINEBELOW", (0, 0), (-1, -1), .3, PDF_LINE), ("BOTTOMPADDING", (0, 0), (-1, -1), 7), ("TOPPADDING", (0, 0), (-1, -1), 7)]))
            block.append(compact_table)
            height = sum(flow.wrap(section_width, 100000)[1] for flow in block)
            if height > 580:
                flush_panels()
                story.extend(block)
            else:
                pending.append(block)
                if len(pending) == 2:
                    flush_panels()
            continue
        rest = [key for key in keys if key != identity]
        graphics = []
        if name == "Trend":
            currencies = ["RNB"] if filters.get("floor_id") == "third-floor" else ["INR"]
            if any("sales_rnb" in row and numeric(row.get("sales_rnb")) and row["sales_rnb"] != 0 for row in rows) and "RNB" not in currencies:
                currencies.append("RNB")
            graphics = [graphic for currency in currencies if (graphic := trend(rows, currency, (section_width - 12) / len(currencies), "Collections over time" if filters.get("date_basis") == "payment" or metadata.get("Date basis") == "payment" else "Sales over time"))]
        elif name not in {"Source Details", "Orders", "Records"}:
            for field, currency in (("sales", "INR"), ("sales_rnb", "RNB"), ("sales", "RNB")):
                if any(numeric(row.get(field)) for row in rows):
                    graphic = ranking(rows, field, currency, (section_width - 12) / 2, name)
                    if graphic:
                        graphics.append(graphic)
            if not graphics and any(numeric(row.get("customers")) for row in rows):
                graphic = ranking(rows, "customers", None, (section_width - 12) / 2, name)
                if graphic:
                    graphics.append(graphic)
        if graphics:
            chart_grid = Table([graphics[i:i + 2] for i in range(0, len(graphics), 2)], colWidths=[section_width / min(2, len(graphics))] * min(2, len(graphics)), hAlign="LEFT")
            chart_grid.setStyle(TableStyle(table_style + [("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
            story.append(chart_grid)
        bands = [[identity, *rest[offset:offset + 5]] for offset in range(0, len(rest), 5)] or [[identity]]
        scales = {}
        for row in rows:
            for key, value in row.items():
                if numeric(value) and key in PDF_MONEY_FIELDS | {"orders", "customers", "quantity"}:
                    ledger = "RNB" if key.endswith("_rnb") else row.get("currency", "INR")
                    scale_key = (key, ledger, str(row.get("unit")) if key == "quantity" else None)
                    lower, upper = scales.get(scale_key, (0, 0))
                    scales[scale_key] = min(lower, value), max(upper, value)
        for index, band in enumerate(bands):
            if len(bands) > 1:
                story.append(para(f"{name} / detail columns {index + 1} of {len(bands)} / {len(rows)} recorded rows", styles["DetailCaption"]))
            widths = [section_width / len(band)] * len(band)
            if name in {"Source Details", "Orders", "Records"} and len(band) > 1:
                widths = [section_width * .46] + [section_width * .54 / (len(band) - 1)] * (len(band) - 1)
            values = [[para("Orders (all selected currencies)" if name == "Trend" and key == "orders" else _label(key), styles["NumberReport"] if any(isinstance(row.get(key), (int, float)) and not isinstance(row.get(key), bool) for row in rows) else body) for key in band]]
            for row in rows:
                cells = []
                for key in band:
                    effective = row.get("metric", "").replace(" ", "_") if key == "value" else key
                    value = _label(row.get("metric", "").replace(" ", "_")) if key == "metric" and name in {"Summary", "Furniture ledger"} else row.get(key)
                    is_number = numeric(value)
                    content = para(display(value, effective, row.get("currency", "INR")), styles["NumberReport"] if is_number else body)
                    if numeric(value) and key in PDF_MONEY_FIELDS | {"orders", "customers", "quantity"}:
                        currency = "RNB" if key.endswith("_rnb") else row.get("currency", "INR")
                        unit = row.get("unit") if key == "quantity" else None
                        lower, upper = scales[(key, currency, str(unit) if key == "quantity" else None)]
                        content = [content, cell_bar(value, lower, upper, widths[band.index(key)] - 12)]
                    cells.append(content)
                values.append(cells)
            table = Table(values, repeatRows=1, colWidths=widths, hAlign="LEFT", splitInRow=1)
            table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), PDF_WASH), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEBELOW", (0, 0), (-1, 0), .6, PDF_LINE), ("LINEBELOW", (0, 1), (-1, -1), .3, PDF_LINE), ("BOTTOMPADDING", (0, 0), (-1, -1), 4), ("TOPPADDING", (0, 0), (-1, -1), 4), ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7)]))
            story.extend([table, Spacer(1, 5)])
    flush_panels()
    if empty_sections and not any(rows for _, rows in sections):
        story.append(para("No matching recorded data: " + ", ".join(empty_sections) + ".", small))
    notes = report.get("coverage", {}).get("notes", report.get("notes", []))
    if notes:
        story.append(para("Data notes", styles["SectionTitle"]))
    for note in notes:
        story.append(para(note, small))
    def page_chrome(canvas, doc):
        canvas.saveState()
        canvas.setTitle(str(title))
        canvas.setAuthor("Buildcon House")
        canvas.setFont(face, 8)
        canvas.setFillColor(PDF_SECONDARY)
        canvas.drawString(36, A4[1] - 25, "Buildcon House · Sales data")
        canvas.setStrokeColor(PDF_LINE)
        canvas.line(36, A4[1] - 34, A4[0] - 36, A4[1] - 34)
        canvas.line(36, 34, A4[0] - 36, 34)
        canvas.setFont(face, 7)
        canvas.drawString(36, 22, "Amounts show their currency. Unavailable is not zero.")
        canvas.drawRightString(A4[0] - 36, 22, f"Page {doc.page}")
        canvas.restoreState()
    SimpleDocTemplate(blob, pagesize=A4, rightMargin=36, leftMargin=36, topMargin=48, bottomMargin=46, allowSplitting=True).build(story, onFirstPage=page_chrome, onLaterPages=page_chrome)
    return blob.getvalue()
