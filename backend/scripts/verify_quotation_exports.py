"""Generate reviewable quotation exports from synthetic data only.

Run from backend/: .venv/bin/python scripts/verify_quotation_exports.py
No application server, database, credentials or network access are required.
"""
from __future__ import annotations

from io import BytesIO
import json
from pathlib import Path
import sys
from time import perf_counter
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from PIL import Image, ImageDraw
from pypdf import PdfReader
from pptx import Presentation

import pdf_generator
import pdf_tiles
from services import quotation_ppt


def fixture_photo(index: int) -> bytes:
    """Portrait, landscape, transparent/hidden-part and diagram test imagery."""
    size = [(600, 1000), (1200, 700), (900, 900), (700, 1100)][index % 4]
    image = Image.new("RGBA", size, (255, 255, 255, 0))
    draw = ImageDraw.Draw(image)
    colour = [(34, 100, 170), (35, 135, 85), (155, 65, 145), (190, 105, 25)][index % 4]
    left, top, right, bottom = size[0] // 5, size[1] // 8, size[0] * 4 // 5, size[1] * 7 // 8
    draw.rounded_rectangle((left, top, right, bottom), radius=30, fill=(*colour, 255))
    draw.text((left + 20, top + 25), f"SYNTHETIC PRODUCT {index + 1}", fill="white", font_size=30)
    draw.ellipse((left + 35, bottom - 130, right - 35, bottom - 35), fill="white")
    stream = BytesIO()
    image.save(stream, format="PNG")
    return stream.getvalue()


def main() -> None:
    output = Path(__file__).resolve().parents[2] / "output" / "exports" / "2026-10-05"
    output.mkdir(parents=True, exist_ok=True)
    images = {f"https://fixture.test/product-{i}.png": fixture_photo(i) for i in range(22)}
    items = [{
        "id": f"line-{i}", "product_id": f"synthetic-{i}", "sku": f"QA-{i + 1:02}",
        "name": "Concealed rough-in body" if i == 16 else f"Synthetic bathroom fitting {i + 1}",
        "room": "Master Bathroom", "area": ["Basin Area", "WC Area", "Shower Area"][i % 3],
        "image": f"https://fixture.test/product-{i}.png", "qty": 1, "unit_price": 123456.78,
        "mrp": 123456.78, "discount_pct": 10, "sort_order": i,
    } for i in range(17)]
    items[4]["seat_cover_image_urls"] = ["https://fixture.test/product-19.png"]
    items[1]["dimension_image_urls"] = ["https://fixture.test/product-20.png", "https://fixture.test/product-21.png"]
    items.append({**items[0], "id": "powder", "sku": "QA-POWDER", "room": "Powder Room", "area": "Basin Area"})
    document = {
        "number": "QA-EXPORT", "customer_name": "Synthetic Export Review", "referrer_name": "Fixture Studio",
        "referrer_type": "architect", "items": items, "rooms": ["Master Bathroom", "Powder Room"],
        "subtotal": 2222222.04, "grand_total": 1999999.84, "doc_date": "05-10-2026",
        "address_snapshot": "Synthetic address for export verification",
    }
    metrics = {}
    with patch.object(quotation_ppt, "remote_image_bytes", side_effect=images.get), \
            patch.object(quotation_ppt, "prefetch_urls"), \
            patch.object(pdf_generator, "_remote_image_bytes", side_effect=images.get), \
            patch.object(pdf_generator, "prefetch_urls"):
        started = perf_counter()
        ppt = quotation_ppt.build_product_ppt(document)
        (output / "one-room-one-page.pptx").write_bytes(ppt)
        deck = Presentation(BytesIO(ppt))
        metrics["ppt"] = {"slides": len(deck.slides), "bytes": len(ppt), "generation_ms": round((perf_counter() - started) * 1000),
                          "product_photos_per_slide": [sum(shape.shape_type == 13 for shape in slide.shapes) for slide in list(deck.slides)[1:]]}
        assert metrics["ppt"]["slides"] == 3
        assert metrics["ppt"]["product_photos_per_slide"] == [18, 1]
        for name, dimensions in [("sanitary", False), ("sanitary-dimensions", True)]:
            data = pdf_generator.build_quotation_pdf({**document, "include_dimensions": dimensions}, {"name": document["customer_name"]})
            (output / f"{name}.pdf").write_bytes(data)
            pages = PdfReader(BytesIO(data)).pages
            text = " ".join(page.extract_text() or "" for page in pages)
            assert all(item["sku"] in text for item in items)
            assert "Concealed rough-in body" in text
            metrics[name] = {"pages": len(pages), "bytes": len(data), "all_skus_present": True}
        tile_items = [{**item, "rate_sqft": 1234.56, "offer_rate": 1200, "rate_box": 12000,
                       "net_amount": 12000, "quantity_unit": "Box", "pcs_per_box": 2, "size": "600 x 1200 mm"} for item in items]
        for name, builder in [("tiles-selection", pdf_tiles.build_tiles_selection_pdf), ("tiles-quotation", pdf_tiles.build_tiles_quotation_pdf)]:
            data = builder({**document, "items": tile_items}, {"name": document["customer_name"]})
            (output / f"{name}.pdf").write_bytes(data)
            pages = PdfReader(BytesIO(data)).pages
            text = " ".join(page.extract_text() or "" for page in pages)
            # Tile forms intentionally print the product detail, without a
            # sanitary article-number column.
            assert all(item["name"] in text for item in items)
            metrics[name] = {"pages": len(pages), "bytes": len(data), "all_products_present": True}
    (output / "verification.json").write_text(json.dumps(metrics, indent=2))
    print(json.dumps({"output": str(output), "verification": metrics}, indent=2))


if __name__ == "__main__":
    main()
