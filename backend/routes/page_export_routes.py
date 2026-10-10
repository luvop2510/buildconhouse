"""Turn the current staff page view into a safe Excel workbook.

The client sends only text already rendered for the signed-in user. This route
does not query business collections or expand that user's data access.
"""
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from pydantic import BaseModel, Field

from auth import get_current_user
from models import UserPublic
from db import db
from services.access_grants import grants_allow, grants_for_user
from services.export import spreadsheet_cell

router = APIRouter(prefix="/exports", tags=["exports"])


class ExportSection(BaseModel):
    name: str = Field(max_length=80)
    rows: list[list[str]] = Field(max_length=5000)


class PageWorkbook(BaseModel):
    page: str = Field(max_length=80)
    resource: str | None = Field(default=None, max_length=80)
    sections: list[ExportSection] = Field(min_length=1, max_length=30)


@router.post("/page.xlsx")
async def export_page_workbook(body: PageWorkbook, user: UserPublic = Depends(get_current_user)):
    if user.custom_access:
        if not user.active_floor_id or not body.resource or not grants_allow(
            await grants_for_user(db, user.id), user_id=user.id,
            resource=body.resource, action="export", floor_id=user.active_floor_id,
        ):
            raise HTTPException(status_code=403, detail="Export permission is required for this page and floor")
    if sum(len(section.rows) for section in body.sections) > 10_000:
        raise HTTPException(status_code=413, detail="Page export is too large; choose one section or brand")
    wb = Workbook()
    wb.remove(wb.active)
    for index, section in enumerate(body.sections, start=1):
        title = section.name.strip() or f"Section {index}"
        for character in "[]:*?/\\":
            title = title.replace(character, " ")
        ws = wb.create_sheet(title[:27] + (f" {index}" if index > 1 else ""))
        ws.append([spreadsheet_cell(body.page[:200]), spreadsheet_cell(section.name[:200])])
        ws.append([])
        for row in section.rows:
            ws.append([spreadsheet_cell(str(cell)[:2000]) for cell in row[:30]])
        ws.freeze_panes = "A3"
        ws.column_dimensions["A"].width = 36
        for column in "BCDEFGHIJ":
            ws.column_dimensions[column].width = 22
    blob = BytesIO()
    wb.save(blob)
    blob.seek(0)
    return StreamingResponse(blob, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={
        "Content-Disposition": 'attachment; filename="buildcon-page-export.xlsx"',
        "Cache-Control": "no-store",
    })
