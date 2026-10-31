"""Follow-ups · Sales Command Center — the API surface.

Reuses existing collections/helpers everywhere possible:
  * payment aggregation  -> routes.payment_routes._paid_by_quotation / ORDER_STATUSES / _clean_phone
  * timeline             -> services.activity_log.timeline_for
  * scoring / reconcile   -> services.followup_engine

No new business logic is duplicated — this module is orchestration + reads.
"""
from __future__ import annotations
from services.followup_engine import reconcile_followups_admitted

from services.export import spreadsheet_cell
from services.multi_filters import predicate, values
from services.referrer_identity import acquired_counts, legacy_referrer_name

import asyncio
import base64
import binascii
import csv
import hashlib
import io
import json
import re
from collections import Counter
from datetime import datetime, timedelta, timezone
from typing import Optional
from uuid import uuid4, uuid5, NAMESPACE_URL

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel
from pymongo.errors import DuplicateKeyError

from services.referrer_identity import referrer_query
from auth import floor_for_write, floor_inherit, floor_query, floor_scope_ids, get_current_user, get_floor_scoped_or_404, require_floor_access, require_min_role, require_roles, accessible_floor_ids
from db import db
from models import (
    AutomationRuleUpdate, Followup, FollowupCallOutcomePayload, FollowupCompletePayload,
    FollowupContactPayload, FollowupCreate, FollowupInterestsPayload, FollowupSavedView,
    FollowupSavedViewCreate, FollowupSnoozePayload, FollowupUpdate,
    NotebookCellPatchPayload, NotebookConversionPayload, NotebookFollowupCreatePayload,
    NotebookOutcomePayload, NotebookReferrerPayload,
    UserPublic, now_iso,
)
from services import automation_rules, workflow_transitions
from services import lead_scope
from services.activity_log import log_event, timeline_for
from services.followup_engine import (
    RULE_DEFINITIONS, _followup_sort_key, age_days,
    compute_bucket, ist_day_bounds_utc, money_short, parse_iso,
    reason_factors_for, reconcile_followups, score_followup,
)
from services.followup_notebook import (
    NotebookConflictError, NotebookValidationError, complete_notebook_row, convert_notebook_row,
    notebook_query, notebook_search_query, normalize_mobile, patch_notebook_row,
    serialize_notebook_row, timeline_event_for_field, validate_notebook_patch, resolve_or_create_customer,
)

router = APIRouter(prefix="/followups", tags=["followups"])
NOTEBOOK_FLOOR_IDS = frozenset({"second-floor", "third-floor"})


# ─────────────────────────────────────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────────────────────────────────────
def _required_action_note(notes: str | None) -> str:
    note = notes.strip() if isinstance(notes, str) else ""
    if not note:
        raise HTTPException(status_code=422, detail="A reason is required when marking a client as lost")
    return note


def _reject_action_note(notes: str | None) -> None:
    if notes is not None:
        raise HTTPException(status_code=422, detail="Notes can only be written when marking a client as lost")


def require_notebook_floor(floor_id: str, user: UserPublic) -> None:
    """Limit the notebook surface to its two declared floor modules."""
    require_floor_access(floor_id, user)
    if floor_id not in NOTEBOOK_FLOOR_IDS:
        raise HTTPException(status_code=404, detail="Notebook is available only on Kitchen or Furniture Floor")
    if user.active_floor_id and floor_id != user.active_floor_id:
        raise HTTPException(status_code=404, detail="Follow-up not found")


async def _selected_followup_or_404(followup_id: str, user: UserPublic) -> dict:
    """Record links must stay inside the same floor as the visible inbox."""
    row = await get_floor_scoped_or_404(
        db.followups, followup_id, user, not_found="Follow-up not found", projection={"_id": 0},
    )
    if row["floor_id"] not in floor_scope_ids(user):
        raise HTTPException(status_code=404, detail="Follow-up not found")
    return row


async def _wake_snoozed(user: UserPublic) -> None:
    """Self-healing: any snooze whose timer has elapsed flips back to open so
    it resurfaces in the inbox without needing a background job.

    This update is triggered by a caller's read. Keep it in that caller's
    floor scope: a Ground Floor request must never mutate an overdue snooze
    belonging to another operational floor.
    """
    await db.followups.update_many(
        floor_query(user, {"status": "snoozed", "snoozed_until": {"$lte": now_iso()}}),
        {"$set": {"status": "open", "updated_at": now_iso()}},
    )


async def _all_with_bucket(user: UserPublic | None = None) -> list[dict]:
    sales = {"category": {"$nin": list(lead_scope.POST_ORDER_CATEGORIES)}, "rule_type": {"$nin": list(lead_scope.POST_ORDER_RULES)}}
    docs = await db.followups.find(floor_query(user, sales) if user else sales, {"_id": 0}).to_list(10000)
    for d in docs:
        d["bucket"] = compute_bucket(d)
        d["effective_priority_level"] = d.get("manual_priority_override") or d.get("priority_level")
    return docs


async def _rule_counts(user: UserPublic) -> dict[str, int]:
    pipeline = [
        {"$match": floor_query(user, {"status": {"$in": ["open", "snoozed"]}})},
        {"$group": {"_id": "$rule_type", "count": {"$sum": 1}}},
    ]
    rows = await db.followups.aggregate(pipeline).to_list(30)
    return {r["_id"]: r["count"] for r in rows}


def _get(d: dict, key: str, default=None):
    return d.get(key, default)


def _encode_notebook_cursor(updated_at: str, row_id: str) -> str:
    raw = json.dumps({"updated_at": updated_at, "id": row_id}, separators=(",", ":")).encode()
    return base64.urlsafe_b64encode(raw).decode()


def _decode_notebook_cursor(cursor: str) -> dict:
    try:
        value = json.loads(base64.urlsafe_b64decode(cursor.encode()).decode())
        if not value.get("updated_at") or not value.get("id"):
            raise ValueError
        return value
    except (ValueError, TypeError, binascii.Error, json.JSONDecodeError):
        raise HTTPException(status_code=400, detail="Invalid notebook cursor")


def _notebook_conflict(error: NotebookConflictError) -> HTTPException:
    return HTTPException(status_code=409, detail={
        "message": str(error), "row": error.row, "changed_fields": error.changed_fields,
    })


async def _notebook_referrer_fields(referrer_id: str | None, floor_id: str) -> dict:
    """Resolve a notebook partner from the directory pinned to its floor.

    The URL floor must match the selected workspace, as enforced by
    ``require_notebook_floor``. Keep the related lookup pinned to that same
    validated floor so another floor's partner cannot enter this row.
    """
    if not referrer_id:
        return {"referrer_id": None, "referrer_name": None, "referrer_type": None}
    referrer = await db.referrers.find_one(referrer_query({"floor_id": floor_id}, {"id": referrer_id}), {"_id": 0})
    if not referrer:
        raise HTTPException(status_code=404, detail="Referrer not found")
    if referrer.get("active") is False:
        raise HTTPException(status_code=400, detail="This referrer is archived")
    return {
        "referrer_id": referrer["id"], "referrer_name": referrer["name"],
        "referrer_type": referrer["type"],
    }


# ─────────────────────────────────────────────────────────────────────────────
# Automation
# ─────────────────────────────────────────────────────────────────────────────
@router.post("/reconcile")
async def reconcile(_: UserPublic = Depends(require_min_role("manager"))):
    return await reconcile_followups()


@router.get("/config/rules")
async def rules_config(user: UserPublic = Depends(get_current_user)):
    counts = await _rule_counts(user)
    return [{**r, "active_count": counts.get(r["rule_type"], 0)} for r in RULE_DEFINITIONS]


@router.get("/config/automation-rules")
async def get_automation_rules(user: UserPublic = Depends(get_current_user)):
    """Configurable reminder cadences (day thresholds) per workspace category
    — see services/automation_rules.py. Editable by manager+ only."""
    return [r for r in await automation_rules.list_rules() if r.get("category") not in lead_scope.POST_ORDER_CATEGORIES]


@router.put("/config/automation-rules/{category}")
async def put_automation_rule(category: str, body: AutomationRuleUpdate, user: UserPublic = Depends(require_min_role("manager"))):
    if category in lead_scope.POST_ORDER_CATEGORIES:
        raise HTTPException(422, "Follow-ups cover only activities before an order is placed")
    patch = body.model_dump(exclude_unset=True)
    updated = await automation_rules.update_rule(category, patch, user_id=user.id, user_name=user.full_name)
    asyncio.create_task(reconcile_followups_admitted())
    return updated


@router.get("/config/assignees")
async def assignees(_: UserPublic = Depends(get_current_user)):
    return await db.users.find(
        {"active": True, "role": {"$ne": "owner"}}, {"_id": 0, "id": 1, "full_name": 1, "role": 1},
    ).sort("full_name", 1).to_list(100)


class WorkflowTransitionUpdate(BaseModel):
    title: Optional[str] = None
    message_template: Optional[str] = None
    priority: Optional[str] = None
    is_active: Optional[bool] = None


@router.get("/config/workflow-transitions")
async def get_workflow_transitions(_: UserPublic = Depends(get_current_user)):
    """Configurable business-event → operational-follow-up transitions (2026-08
    CRM foundation) — e.g. Quotation → Order Confirmed. See
    services/workflow_transitions.py / followup_engine._operational_followup_producer.
    Editable by manager+ only."""
    return [r for r in await workflow_transitions.list_transitions() if r.get("key") != "order_confirmed"]


@router.put("/config/workflow-transitions/{key}")
async def put_workflow_transition(key: str, body: WorkflowTransitionUpdate, user: UserPublic = Depends(require_min_role("manager"))):
    if key == "order_confirmed":
        raise HTTPException(422, "Order activity belongs in Orders")
    patch = body.model_dump(exclude_unset=True)
    updated = await workflow_transitions.update_transition(key, patch, user_id=user.id, user_name=user.full_name)
    asyncio.create_task(reconcile_followups_admitted())
    return updated


# ─────────────────────────────────────────────────────────────────────────────
# KPIs / Today's Mission / Insights — all literal paths, MUST precede /{id}
# ─────────────────────────────────────────────────────────────────────────────
@router.get("/stats")
async def stats(user: UserPublic = Depends(get_current_user)):
    await _wake_snoozed(user)
    docs = await _all_with_bucket(user)

    counts = {b: 0 for b in ("overdue", "today", "tomorrow", "this_week", "later", "completed", "snoozed")}
    overdue_critical = 0
    for d in docs:
        counts[d["bucket"]] = counts.get(d["bucket"], 0) + 1
        if d["bucket"] == "overdue" and d["effective_priority_level"] in ("critical", "high"):
            overdue_critical += 1

    today_critical = sum(1 for d in docs if d["bucket"] == "today" and d["effective_priority_level"] == "critical")
    waiting_for_customer = sum(
        1 for d in docs if d.get("status") == "open" and d.get("rule_type") == "payment_partial"
    )

    # Split "overdue" into payment-specific vs generic — fixes the "which
    # payment is overdue" 5-second-scan gap identified in the UX audit.
    overdue_payments_count = sum(1 for d in docs if d.get("status") == "open" and d.get("rule_type") == "payment_overdue")
    overdue_payments_amount = sum(
        d.get("value", 0) for d in docs if d.get("status") == "open" and d.get("rule_type") == "payment_overdue"
    )
    expiring_quotations_count = sum(
        1 for d in docs if d.get("status") == "open" and d.get("rule_type") == "quotation_expiring"
    )

    start, _end = ist_day_bounds_utc(0)
    y_start, y_end = ist_day_bounds_utc(-1)
    completed_today = sum(1 for d in docs if d.get("completed_at") and d["completed_at"] >= start.isoformat())
    completed_yesterday = sum(
        1 for d in docs if d.get("completed_at") and y_start.isoformat() <= d["completed_at"] < y_end.isoformat()
    )

    rules = await _rule_counts(user)
    workspace_counts = {
        "selection": rules.get("selection_waiting", 0),
        "quotation_tiles": rules.get("quotation_tiles_waiting", 0),
        "payment": rules.get("payment_overdue", 0) + rules.get("payment_partial", 0),
        "walk_in": 0,  # Phase 4 — Walk-ins module not yet built
    }
    return {
        "today_tasks": counts["today"], "today_critical": today_critical,
        "overdue": counts["overdue"], "overdue_critical": overdue_critical,
        "overdue_payments_count": overdue_payments_count,
        "overdue_payments_amount": round(overdue_payments_amount, 2),
        "overdue_payments_amount_short": money_short(overdue_payments_amount),
        "expiring_quotations_count": expiring_quotations_count,
        "tomorrow": counts["tomorrow"],
        "this_week": counts["this_week"],
        "waiting_for_customer": waiting_for_customer,
        "completed_today": completed_today,
        "completed_trend": completed_today - completed_yesterday,
        "snoozed": counts["snoozed"],
        "later": counts["later"],
        "rules": [{**r, "active_count": rules.get(r["rule_type"], 0)} for r in RULE_DEFINITIONS],
        "workspace_counts": workspace_counts,
    }


@router.get("/mission")
async def mission(user: UserPublic = Depends(get_current_user)):
    await _wake_snoozed(user)
    docs = await _all_with_bucket(user)
    actionable = [d for d in docs if d["bucket"] in ("overdue", "today")]

    revenue_at_risk = sum(d.get("value", 0) for d in actionable)
    overdue_payments = sum(1 for d in actionable if d["rule_type"] == "payment_overdue")
    expiring_today = sum(1 for d in docs if d["rule_type"] == "quotation_expiring" and d["bucket"] == "today")
    critical_count = sum(1 for d in actionable if d["effective_priority_level"] == "critical")

    minutes = 0
    for d in actionable:
        minutes += 6 if d["suggested_channel"] == "call" else 2 if d["suggested_channel"] == "whatsapp" else 3

    top = sorted(actionable, key=lambda d: -(d.get("priority_score") or 0))[:3]
    return {
        "due_count": len(actionable),
        "revenue_at_risk": round(revenue_at_risk, 2),
        "revenue_at_risk_short": money_short(revenue_at_risk),
        "overdue_payments": overdue_payments,
        "quotations_expiring_today": expiring_today,
        "critical_count": critical_count,
        "estimated_minutes": minutes,
        "top_priorities": [
            {"id": d["id"], "customer_name": d["customer_name"], "reason": d["reason"], "priority_score": d["priority_score"]}
            for d in top
        ],
        "greeting_name": (user.full_name or "").split()[0] if user.full_name else "there",
    }


@router.get("/insights")
async def insights(user: UserPublic = Depends(get_current_user)):
    start, end = ist_day_bounds_utc(0)
    rng = {"$gte": start.isoformat(), "$lt": end.isoformat()}

    calls = await db.activity_events.count_documents(
        floor_query(user, {"event_type": "followup.call_logged", "created_at": rng})
    )
    whatsapps = await db.activity_events.count_documents(floor_query(user, {
        "event_type": "followup.contacted", "payload.channel": "whatsapp", "created_at": rng,
    }))
    quotations_approved = await db.quotations.count_documents(floor_query(user, {
        "status": {"$in": ["approved", "won"]}, "updated_at": rng,
    }))
    completed_today = await db.followups.count_documents(floor_query(user, {"completed_at": rng}))
    still_open = await db.followups.count_documents(
        floor_query(user, {"status": {"$in": ["open", "snoozed"]}})
    )
    response_rate = round(100 * completed_today / max(1, completed_today + still_open))

    return {
        "calls_completed": calls,
        "whatsapps_sent": whatsapps,
        "payments_collected": 0,
        "quotations_approved": quotations_approved,
        "response_rate": min(100, response_rate),
    }


# ─────────────────────────────────────────────────────────────────────────────
# Export — CSV / Excel of the current filtered list. Saved Views — persisted
# filter configurations per user. Both literal-path routes; MUST precede /{id}.
# ─────────────────────────────────────────────────────────────────────────────
@router.get("/export")
async def export_followups(
    origin: Optional[str] = None,
    format: str = Query("xlsx", pattern="^(xlsx|csv)$"),
    bucket: Optional[str] = None,
    priority: Optional[str] = None,
    category: Optional[str] = None,
    customer_tier: Optional[str] = None,
    assigned_to: Optional[str] = None,
    q: Optional[str] = None,
    user: UserPublic = Depends(get_current_user),
):
    rows = await list_followups(
        bucket=bucket, priority=priority, category=category, channel=None,
        customer_tier=customer_tier, assigned_to=assigned_to, q=q, origin=origin, limit=3000, user=user,
    )
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M")

    if format == "csv":
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow([spreadsheet_cell(value) for value in ["Customer", "Phone", "Type", "Reason", "Next Action", "Value", "Priority", "Score", "Due", "Status", "Assigned To"]])
        for d in rows:
            writer.writerow([spreadsheet_cell(value) for value in [
                d.get("customer_name"), d.get("customer_phone"), d.get("category"), d.get("reason"),
                d.get("next_action"), d.get("value"), d.get("effective_priority_level") or d.get("priority_level"),
                d.get("priority_score"), d.get("due_at"), d.get("status"), d.get("assigned_to_name"),
            ]])
        mem = io.BytesIO(buf.getvalue().encode("utf-8"))
        return StreamingResponse(mem, media_type="text/csv", headers={
            "Content-Disposition": f'attachment; filename="followups-{stamp}.csv"',
        })

    wb = Workbook()
    ws = wb.active
    ws.title = "Follow-ups"
    ws["A1"] = "BuildCon House — Follow-ups Export"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:K1")
    ws["A2"] = f"{len(rows)} follow-ups · Exported {datetime.now(timezone.utc).strftime('%d %b %Y · %H:%M UTC')}"
    ws["A2"].font = Font(color="6B7280", size=10)
    ws.merge_cells("A2:K2")

    headers = ["Customer", "Phone", "Type", "Reason", "Next Action", "Value (₹)", "Priority", "Score", "Due", "Status", "Assigned To"]
    header_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col, value=spreadsheet_cell(h))
        cell.font = Font(bold=True, color="374151", size=11)
        cell.fill = header_fill

    for i, d in enumerate(rows, start=5):
        ws.cell(row=i, column=1, value=spreadsheet_cell(d.get("customer_name")))
        ws.cell(row=i, column=2, value=spreadsheet_cell(d.get("customer_phone")))
        ws.cell(row=i, column=3, value=spreadsheet_cell(d.get("category")))
        ws.cell(row=i, column=4, value=spreadsheet_cell(d.get("reason")))
        ws.cell(row=i, column=5, value=spreadsheet_cell(d.get("next_action")))
        ws.cell(row=i, column=6, value=spreadsheet_cell(d.get("value")))
        ws.cell(row=i, column=7, value=(spreadsheet_cell(d.get("effective_priority_level") or d.get("priority_level"))))
        ws.cell(row=i, column=8, value=spreadsheet_cell(d.get("priority_score")))
        ws.cell(row=i, column=9, value=spreadsheet_cell(d.get("due_at")))
        ws.cell(row=i, column=10, value=spreadsheet_cell(d.get("status")))
        ws.cell(row=i, column=11, value=spreadsheet_cell(d.get("assigned_to_name")))

    for i, w in enumerate([22, 16, 12, 44, 22, 12, 10, 8, 22, 12, 18], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={
        "Content-Disposition": f'attachment; filename="followups-{stamp}.xlsx"',
    })


@router.get("/saved-views")
async def list_saved_views(user: UserPublic = Depends(get_current_user)):
    return await db.followup_saved_views.find({"user_id": user.id}, {"_id": 0}).sort("created_at", -1).to_list(50)


@router.post("/saved-views", response_model=FollowupSavedView)
async def create_saved_view(body: FollowupSavedViewCreate, user: UserPublic = Depends(get_current_user)):
    v = FollowupSavedView(user_id=user.id, name=body.name, filters=body.filters)
    await db.followup_saved_views.insert_one(v.model_dump())
    return v


@router.delete("/saved-views/{view_id}")
async def delete_saved_view(view_id: str, user: UserPublic = Depends(get_current_user)):
    await db.followup_saved_views.delete_one({"id": view_id, "user_id": user.id})
    return {"ok": True}


# ─────────────────────────────────────────────────────────────────────────────
# List
# ─────────────────────────────────────────────────────────────────────────────
def _paged_followup_query(
    user: UserPublic,
    *,
    bucket: Optional[str], priority: Optional[str], category: Optional[str],
    customer_tier: Optional[str], assigned_to: Optional[str], q: Optional[str],
    rule_type: Optional[str], view: str, waiting: bool, completed_today: bool,
    origin: Optional[str] = None,
) -> dict:
    """Build the database-side portion of the paged inbox filter.

    The legacy list endpoint intentionally retains its array response for
    integration callers. The canonical UI uses this bounded path so a search
    is not silently limited to the first thousand rows.
    """
    buckets = values(bucket)
    if len(buckets) > 1:
        return {"$or": [_paged_followup_query(
            user, bucket=choice, priority=priority, category=category, customer_tier=customer_tier,
            assigned_to=assigned_to, q=q, rule_type=rule_type, view=view, waiting=waiting,
            completed_today=completed_today, origin=origin,
        ) for choice in buckets]}
    expiring = bucket == "expiring"
    if expiring:
        rule_type, bucket = "quotation_expiring", None
    elif bucket == "waiting":
        waiting, bucket = True, None
    elif bucket == "completed_today":
        completed_today, bucket = True, None
    clauses: list[dict] = [{"category": {"$nin": list(lead_scope.POST_ORDER_CATEGORIES)}, "rule_type": {"$nin": list(lead_scope.POST_ORDER_RULES)}}]
    if origin:
        clauses.append({"origin": predicate(origin)})
    if category:
        clauses.append({"category": predicate(category)})
    if customer_tier:
        clauses.append({"customer_tier": predicate(customer_tier)})
    if assigned_to:
        clauses.append(_assigned_filter(user, assigned_to))
    if rule_type:
        clauses.append({"rule_type": predicate(rule_type)})
    if q:
        term = {"$regex": re.escape(q), "$options": "i"}
        clauses.append({"$or": [
            {"customer_name": term}, {"customer_phone": term}, {"quotation_number": term},
            {"purchase_number": term}, {"project_name": term}, {"reason": term}, {"tags": term},
        ]})
    if priority:
        clauses.append({"$or": [
            {"manual_priority_override": predicate(priority)},
            {"manual_priority_override": None, "priority_level": predicate(priority)},
        ]})

    today_start, tomorrow_start = ist_day_bounds_utc()
    _, day_after_tomorrow = ist_day_bounds_utc(1)
    _, week_end = ist_day_bounds_utc(7)
    today = today_start.isoformat()
    tomorrow = tomorrow_start.isoformat()
    day_after = day_after_tomorrow.isoformat()
    week_ends = week_end.isoformat()
    unresolved = {
        "$and": [
            {"status": {"$nin": ["done", "dismissed"]}},
            {"$or": [{"status": {"$ne": "snoozed"}}, {"snoozed_until": None}]},
        ],
    }
    if bucket and bucket != "all":
        if bucket == "completed":
            clauses.append({"status": {"$in": ["done", "dismissed"]}})
        elif bucket == "snoozed":
            clauses.append({"status": "snoozed", "snoozed_until": {"$ne": None}})
        else:
            ranges = {
                "overdue": {"$lt": today},
                "today": {"$gte": today, "$lt": tomorrow},
                "tomorrow": {"$gte": tomorrow, "$lt": day_after},
                "this_week": {"$gte": day_after, "$lt": week_ends},
                "later": {"$gte": week_ends},
            }
            if bucket in ranges:
                clauses.extend([unresolved, {"due_at": ranges[bucket]}])
    if expiring:
        clauses.append({"status": "open"})
    if waiting:
        clauses.append({"status": "open", "rule_type": "payment_partial"})
    if completed_today:
        clauses.extend([
            {"status": {"$in": ["done", "dismissed"]}},
            {"completed_at": {"$gte": today, "$lt": tomorrow}},
        ])
    if view == "won":
        if not assigned_to:
            clauses.append({"assigned_to": user.id})
        clauses.append({"auto_resolved": True, "completed_outcome": {"$in": ["won", "converted"]}})
    elif view == "lost":
        clauses.append({"completed_outcome": "lost"})
    else:
        clauses.append({"$nor": [
            {"auto_resolved": True, "completed_outcome": {"$in": ["won", "converted"]}},
            {"completed_outcome": "lost"},
        ]})

    base = clauses[0] if len(clauses) == 1 else {"$and": clauses} if clauses else {}
    return floor_query(user, base)


async def _mask_furniture_owner_salesperson(rows):
    """Keep order ownership for filtering, without presenting owners as salespeople."""
    ids = {row.get("assigned_to") for row in rows if row.get("floor_id") == "third-floor" and row.get("assigned_to")}
    if not ids:
        return
    owners = await db.users.find({"id": {"$in": list(ids)}, "role": "owner"}, {"_id": 0, "id": 1}).to_list(len(ids))
    owner_ids = {owner["id"] for owner in owners}
    for row in rows:
        if row.get("floor_id") == "third-floor" and row.get("assigned_to") in owner_ids:
            row["assigned_to_name"] = None


async def _enrich_followup_references(rows, user, customers=None):
    """Hydrate attribution in bounded floor-local batches, never per lead."""
    await _mask_furniture_owner_salesperson(rows)
    by_floor = {}
    for row in rows:
        if row.get("floor_id"):
            by_floor.setdefault(row["floor_id"], []).append(row)
    for floor, local_rows in by_floor.items():
        customer_ids = list(dict.fromkeys(row["customer_id"] for row in local_rows if row.get("customer_id")))
        if customers is None:
            customer_rows = await db.customers.find({"floor_id": floor, "id": {"$in": customer_ids}}, {
                "_id": 0, "id": 1, "referrer_id": 1, "referrer_name": 1, "referrer_type": 1,
                "architect_name": 1, "interior_designer_name": 1, "architect_interior_designer": 1, "architect": 1,
            }).to_list(len(customer_ids)) if customer_ids else []
        else:
            customer_rows = [customer for customer in customers if customer and customer.get("floor_id") == floor]
        customer_map = {customer["id"]: customer for customer in customer_rows}
        for row in local_rows:
            customer = customer_map.get(row.get("customer_id"), {})
            if not row.get("referrer_id") and customer.get("referrer_id"):
                for key in ("referrer_id", "referrer_name", "referrer_type"):
                    row[key] = customer.get(key)
            if not row.get("referrer_name"):
                row["referrer_name"] = legacy_referrer_name(row) or customer.get("referrer_name") or legacy_referrer_name(customer)
        ids = list(dict.fromkeys(row["referrer_id"] for row in local_rows if row.get("referrer_id")))
        if not ids:
            continue
        selector = {"id": {"$in": ids}}
        permitted = accessible_floor_ids(user)
        if permitted is not None:
            selector["floor_id"] = {"$in": permitted}
        partners, counts = await asyncio.gather(
            db.referrers.find(referrer_query(selector), {"_id": 0, "id": 1, "name": 1, "type": 1, "company": 1}).to_list(len(ids)),
            acquired_counts(db, ids, [floor]),
        )
        directory = {partner["id"]: partner for partner in partners}
        for row in local_rows:
            partner_id = row.get("referrer_id")
            if not partner_id:
                continue
            partner = directory.get(partner_id, {})
            row["referrer_name"] = partner.get("name") or row.get("referrer_name")
            row["referrer_type"] = partner.get("type") or row.get("referrer_type")
            row["referrer_company"] = partner.get("company")
            row["referrer_customers_acquired"] = counts.get(partner_id, 0)
    return rows


@router.get("/page")
async def list_followups_page(
    bucket: Optional[str] = None,
    priority: Optional[str] = None,
    category: Optional[str] = None,
    customer_tier: Optional[str] = None,
    assigned_to: Optional[str] = None,
    q: Optional[str] = None,
    rule_type: Optional[str] = None,
    origin: Optional[str] = None,
    view: str = Query("active", pattern="^(active|won|lost)$"),
    waiting: bool = False,
    completed_today: bool = False,
    page: int = Query(1, ge=1),
    page_size: int = Query(100, ge=25, le=200),
    user: UserPublic = Depends(get_current_user),
):
    """Return one stable, floor-scoped follow-up page for the inbox."""
    await _wake_snoozed(user)
    start_of_today, _ = ist_day_bounds_utc()
    pipeline = [
        {"$match": _paged_followup_query(
            user, bucket=bucket, priority=priority, category=category,
            customer_tier=customer_tier, assigned_to=assigned_to, q=q,
            rule_type=rule_type, origin=origin, view=view, waiting=waiting,
            completed_today=completed_today,
        )},
        {"$addFields": {
            "_inbox_overdue": {"$cond": [{"$and": [
                {"$eq": ["$status", "open"]}, {"$lt": ["$due_at", start_of_today.isoformat()]},
            ]}, 0, 1]},
            "_inbox_unresolved": {"$cond": [{"$in": ["$status", ["open", "snoozed"]]}, 0, 1]},
            "_inbox_mine": {"$cond": [{"$or": [{"$eq": ["$assigned_to", user.id]}, {"$gt": [{"$size": {"$filter": {"input": {"$ifNull": ["$assignments", []]}, "as": "a", "cond": {"$and": [{"$eq": ["$$a.user_id", user.id]}, {"$eq": ["$$a.floor_id", "$floor_id"]}]}}}}, 0]}]}, 0, 1]},
        }},
        {"$sort": {
            "_inbox_overdue": 1, "_inbox_unresolved": 1, "_inbox_mine": 1,
            "priority_score": -1, "due_at": 1, "id": 1,
        }},
        {"$skip": (page - 1) * page_size},
        {"$limit": page_size + 1},
        {"$project": {"_id": 0, "_inbox_overdue": 0, "_inbox_unresolved": 0, "_inbox_mine": 0}},
    ]
    rows = await db.followups.aggregate(pipeline).to_list(page_size + 1)
    has_more = len(rows) > page_size
    rows = rows[:page_size]
    for row in rows:
        row["bucket"] = compute_bucket(row)
        row["effective_priority_level"] = row.get("manual_priority_override") or row.get("priority_level")
    await _enrich_followup_references(rows, user)
    return {"items": rows, "page": page, "page_size": page_size, "has_more": has_more}


@router.get("")
async def list_followups(
    bucket: Optional[str] = None,
    priority: Optional[str] = None,
    category: Optional[str] = None,
    channel: Optional[str] = None,
    customer_tier: Optional[str] = None,
    assigned_to: Optional[str] = None,
    q: Optional[str] = None,
    origin: Optional[str] = None,
    limit: int = Query(1000, ge=1, le=3000),
    user: UserPublic = Depends(get_current_user),
):
    await _wake_snoozed(user)
    query: dict = {"$nor": [{"category": {"$in": list(lead_scope.POST_ORDER_CATEGORIES)}}, {"rule_type": {"$in": list(lead_scope.POST_ORDER_RULES)}}]}
    if isinstance(origin, str) and origin:
        query["origin"] = predicate(origin)
    if category:
        query["category"] = predicate(category)
    if channel:
        query["suggested_channel"] = predicate(channel)
    if customer_tier:
        query["customer_tier"] = predicate(customer_tier)
    if assigned_to:
        query["$and"] = [_assigned_filter(user, assigned_to)]
    if q:
        term = {"$regex": re.escape(q), "$options": "i"}
        query["$or"] = [
            {"customer_name": term}, {"customer_phone": term}, {"quotation_number": term},
            {"purchase_number": term}, {"project_name": term}, {"reason": term}, {"tags": term},
        ]

    docs = await db.followups.find(floor_query(user, query), {"_id": 0}).to_list(limit * 3)
    for d in docs:
        d["bucket"] = compute_bucket(d)
        d["effective_priority_level"] = d.get("manual_priority_override") or d.get("priority_level")

    if bucket and bucket != "all":
        selected_buckets = values(bucket)
        today, tomorrow = ist_day_bounds_utc()
        def matches_bucket(row):
            return any(
                (row.get("status") == "open" and row.get("rule_type") == "quotation_expiring") if choice == "expiring"
                else (row.get("status") == "open" and row.get("rule_type") == "payment_partial") if choice == "waiting"
                else (row.get("status") in {"done", "dismissed"} and today.isoformat() <= (row.get("completed_at") or "") < tomorrow.isoformat()) if choice == "completed_today"
                else row["bucket"] == choice
                for choice in selected_buckets
            )
        docs = [d for d in docs if matches_bucket(d)]
    if priority:
        docs = [d for d in docs if d["effective_priority_level"] in values(priority)]

    docs.sort(key=lambda d: _followup_sort_key(d, user.id))
    return await _enrich_followup_references(docs[:limit], user)


# ─────────────────────────────────────────────────────────────────────────────
# Kitchen/Furniture digital notebook — same followups collection, narrow DTO.
# Literal notebook paths precede /{followup_id} below.
# ─────────────────────────────────────────────────────────────────────────────
@router.get("/notebook/{floor_id}")
async def list_notebook(
    floor_id: str,
    view: str = Query("followups", pattern="^(followups|quotation)$"),
    status: Optional[str] = None,
    q: Optional[str] = None,
    cursor: Optional[str] = None,
    limit: int = Query(100, ge=1, le=250),
    user: UserPublic = Depends(get_current_user),
):
    require_notebook_floor(floor_id, user)
    query = notebook_query(user, floor_id, {"is_converted": view == "quotation"})
    if status and status != "all":
        query["notebook_status"] = predicate(status, allowed={"new", "pending", "won", "lost"})
    search = notebook_search_query(q or "")
    if search:
        query["$or"] = search["$or"]
    if cursor:
        after = _decode_notebook_cursor(cursor)
        query = {"$and": [query, {"$or": [
            {"updated_at": {"$lt": after["updated_at"]}},
            {"updated_at": after["updated_at"], "id": {"$lt": after["id"]}},
        ]}]}
    projection = {
        "_id": 0, **({
            "id": 1, "customer_name": 1, "customer_phone": 1, "address": 1,
            "kitchen_type": 1, "referred_by": 1, "architect_interior_designer": 1,
            "referrer_id": 1, "referrer_name": 1, "referrer_type": 1,
            "interested_floor_ids": 1, "assignments": 1, "reference_source": 1, "floor_id": 1, "lead_group_id": 1,
            "notebook_status": 1, "notes": 1, "is_converted": 1, "updated_at": 1,
            "quotation_price": 1, "estimated_value": 1, "quotation_date": 1,
            "converted_at": 1, "lost_reason": 1, "last_contacted_at": 1, "contact_attempts": 1,
        }),
    }
    rows = await db.followups.find(query, projection).sort([("updated_at", -1), ("id", -1)]).limit(limit + 1).to_list(limit + 1)
    has_more = len(rows) > limit
    rows = rows[:limit]
    return {
        "rows": [serialize_notebook_row(row) for row in rows],
        "next_cursor": _encode_notebook_cursor(rows[-1]["updated_at"], rows[-1]["id"]) if has_more and rows else None,
    }


@router.post("/notebook/{floor_id}")
async def create_notebook_row(
    floor_id: str,
    body: NotebookFollowupCreatePayload,
    user: UserPublic = Depends(get_current_user),
):
    require_notebook_floor(floor_id, user)
    _reject_action_note(body.notes)
    _required_create_assignments(body.interested_floor_ids or [floor_id], body.assignments)
    patch = body.model_dump(exclude_none=True, exclude={"notes", "interested_floor_ids", "assignments", "reference_source"})
    if body.interested_floor_ids or body.assignments is not None:
        await lead_scope.validate_scope(db, user, body.interested_floor_ids or [floor_id], [], source_floor=floor_id)
    requested_referrer_id = patch.pop("referrer_id", None)
    patch["notebook_status"] = "new"
    try:
        clean = validate_notebook_patch(
            patch, converted=False, current={}, creating=True, floor_id=floor_id,
        )
    except NotebookValidationError as error:
        raise HTTPException(status_code=422, detail=str(error))
    customer = await resolve_or_create_customer(
        db, user=user, floor_id=floor_id, name=clean["customer_name"],
        phone=clean["customer_phone"], address=clean.get("address"),
    )
    referrer_fields = await _notebook_referrer_fields(requested_referrer_id, floor_id)
    # A customer has one notebook row per floor. Repeated creation requests
    # must resolve to that row instead of making the register appear to ignore
    # edits or creating a second customer conversation.
    notebook_key = f"{floor_id}:{customer['id']}"
    existing = await db.followups.find_one(
        notebook_query(user, floor_id, {"notebook_key": notebook_key}), {"_id": 0},
    )
    if getattr(user, "custom_access", False) and existing and existing.get("is_converted"):
        raise HTTPException(status_code=403, detail="This customer is already in Quotation Follow-up; open that page with its assigned access")
    if getattr(user, "custom_access", False) and existing and requested_referrer_id:
        raise HTTPException(status_code=409, detail="This walk-in already exists; edit its partner from the existing row")
    if requested_referrer_id:
        await db.customers.update_one(
            {"id": customer["id"], "floor_id": floor_id}, {"$set": referrer_fields},
        )
        customer = {**customer, **referrer_fields}
    if existing:
        # Retrying an add with a newly selected directory partner should not
        # make the picker look as if it saved only on the customer. Keep the
        # already-canonical lead row in sync too.
        if requested_referrer_id:
            await db.followups.update_one(
                notebook_query(user, floor_id, {"notebook_key": notebook_key}),
                {"$set": {**referrer_fields, "architect_interior_designer": referrer_fields["referrer_name"], "updated_at": now_iso()}},
            )
            existing = await db.followups.find_one(
                notebook_query(user, floor_id, {"notebook_key": notebook_key}), {"_id": 0},
            )
        if body.interested_floor_ids or body.assignments is not None or body.reference_source:
            existing = await lead_scope.update_interests(db, user=user, source=existing, floor_ids=body.interested_floor_ids or [floor_id], assignments=body.assignments, metadata={"reference_source": body.reference_source, **referrer_fields})
        return serialize_notebook_row(existing)
    now = now_iso()
    row = Followup(
        floor_id=floor_id, notebook_key=notebook_key, source_key=f"notebook:{notebook_key}",
        rule_type="manual", category="sales", customer_id=customer["id"],
        customer_name=clean["customer_name"], customer_phone=clean["customer_phone"],
        customer_tier=customer.get("tier", "retail"), reason="Notebook follow-up",
        next_action="Call customer", next_action_reason="Notebook customer follow-up",
        suggested_channel="call", priority_score=0, priority_level="medium", due_at=now,
        is_automated=False, is_converted=False, notebook_status="new",
        address=clean.get("address"), kitchen_type=clean.get("kitchen_type"),
        referred_by=clean.get("referred_by"),
        **referrer_fields,
        architect_interior_designer=clean.get("architect_interior_designer"),
        notes=clean.get("notes"),
    )
    if body.interested_floor_ids or body.assignments is not None or body.reference_source:
        await lead_scope.update_interests(db, user=user, source=row.model_dump(), floor_ids=body.interested_floor_ids or [floor_id], assignments=body.assignments, metadata={"reference_source": body.reference_source, **referrer_fields}, dry_run=True)
    try:
        await db.followups.insert_one(row.model_dump())
    except DuplicateKeyError:
        # The unique notebook_key index wins races between two staff members
        # creating the same floor/customer record at once.
        existing = await db.followups.find_one(
            notebook_query(user, floor_id, {"notebook_key": notebook_key}), {"_id": 0},
        )
        if existing:
            if getattr(user, "custom_access", False) and existing.get("is_converted"):
                raise HTTPException(status_code=403, detail="This customer is already in Quotation Follow-up; open that page with its assigned access")
            return serialize_notebook_row(existing)
        raise
    await log_event(
        event_type="project_followup.created", entity_type="followup", entity_id=row.id,
        actor=user, customer_id=row.customer_id, floor_id=floor_id,
        summary="Notebook follow-up created",
    )
    document = row.model_dump()
    if body.interested_floor_ids or body.assignments is not None or body.reference_source:
        document = await lead_scope.update_interests(db, user=user, source=document, floor_ids=body.interested_floor_ids or [floor_id], assignments=body.assignments, metadata={"reference_source": body.reference_source, **referrer_fields})
    return serialize_notebook_row(document)


@router.get("/notebook/{floor_id}/{row_id}")
async def get_notebook_row(floor_id: str, row_id: str, user: UserPublic = Depends(get_current_user)):
    require_notebook_floor(floor_id, user)
    row = await db.followups.find_one(notebook_query(user, floor_id, {"id": row_id}), {"_id": 0})
    if not row:
        raise HTTPException(status_code=404, detail="Notebook row not found")
    return serialize_notebook_row(row)


@router.patch("/notebook/{floor_id}/{row_id}")
async def patch_notebook(
    floor_id: str, row_id: str, body: NotebookCellPatchPayload,
    user: UserPublic = Depends(get_current_user),
):
    require_notebook_floor(floor_id, user)
    field = "notebook_status" if body.field == "status" else body.field
    _reject_action_note(body.notes)
    if field == "notes":
        raise HTTPException(status_code=422, detail="Use Mark as lost to record a loss reason")
    if field == "notebook_status" and isinstance(body.value, str) and body.value in {"won", "lost"}:
        raise HTTPException(status_code=422, detail="Use the explicit Won or Lost action so this outcome is recorded correctly")
    if field == "customer_phone" and body.value is not None:
        value = normalize_mobile(str(body.value))
    else:
        value = body.value
    before = await db.followups.find_one(notebook_query(user, floor_id, {"id": row_id}), {"_id": 0})
    if not before:
        raise HTTPException(status_code=404, detail="Notebook row not found")
    try:
        row = await patch_notebook_row(
            db, user=user, floor_id=floor_id, row_id=row_id,
            patch={field: value}, expected_updated_at=body.updated_at,
        )
    except NotebookConflictError as error:
        raise _notebook_conflict(error)
    except NotebookValidationError as error:
        raise HTTPException(status_code=422, detail=str(error))
    event_type, summary = timeline_event_for_field(field, before.get(field), value)
    if field == "notebook_status" and value == "won":
        event_type, summary = "project_followup.won", "Won"
    elif field == "notebook_status" and value == "lost":
        event_type, summary = "project_followup.lost", "Lost"
    await log_event(
        event_type=event_type, entity_type="followup", entity_id=row_id, actor=user,
        customer_id=before.get("customer_id"), floor_id=floor_id, payload={"field": field, "value": value},
        summary=summary,
    )
    if field == "notebook_status" and value == "lost":
        await log_event(
            event_type="project_followup.lost_note", entity_type="followup", entity_id=row_id, actor=user,
            customer_id=before.get("customer_id"), floor_id=floor_id,
            payload={"note": row.get("notes")}, summary="Lost note recorded",
        )
    return row


@router.post("/notebook/{floor_id}/{row_id}/outcome")
async def complete_notebook_outcome(
    floor_id: str, row_id: str, body: NotebookOutcomePayload,
    user: UserPublic = Depends(get_current_user),
):
    """Close a Kitchen/Furniture lead with a durable Won/Lost outcome."""
    require_notebook_floor(floor_id, user)
    notes = _required_action_note(body.notes or body.lost_reason) if body.outcome == "lost" else None
    if body.outcome != "lost":
        _reject_action_note(body.notes)
    before = await db.followups.find_one(notebook_query(user, floor_id, {"id": row_id}), {"_id": 0})
    if not before:
        raise HTTPException(status_code=404, detail="Notebook row not found")
    try:
        row = await complete_notebook_row(
            db, user=user, floor_id=floor_id, row_id=row_id, outcome=body.outcome,
            lost_reason=notes, expected_updated_at=body.updated_at, notes=notes,
        )
    except NotebookConflictError as error:
        raise _notebook_conflict(error)
    except NotebookValidationError as error:
        raise HTTPException(status_code=422, detail=str(error))
    summary = "Client won" if body.outcome == "won" else f"Client lost — {notes}"
    await log_event(
        event_type=f"project_followup.{body.outcome}", entity_type="followup", entity_id=row_id,
        actor=user, customer_id=before.get("customer_id"), floor_id=floor_id,
        payload={"notes": notes, "lost_reason": body.lost_reason}, summary=summary,
    )
    return row


@router.put("/notebook/{floor_id}/{row_id}/referrer")
async def assign_notebook_referrer(
    floor_id: str, row_id: str, body: NotebookReferrerPayload,
    user: UserPublic = Depends(get_current_user),
):
    """Assign/clear a directory partner on a lead and its linked customer."""
    require_notebook_floor(floor_id, user)
    _reject_action_note(body.notes)
    query = notebook_query(user, floor_id, {"id": row_id})
    before = await db.followups.find_one(query, {"_id": 0})
    if not before:
        raise HTTPException(status_code=404, detail="Notebook row not found")
    if before.get("updated_at") != body.updated_at:
        raise _notebook_conflict(NotebookConflictError(serialize_notebook_row(before), ["referrer_id"]))
    fields = await _notebook_referrer_fields(body.referrer_id, floor_id)
    now = now_iso()
    result = await db.followups.update_one(
        {**query, "updated_at": body.updated_at},
        {"$set": {**fields, "architect_interior_designer": fields["referrer_name"] or "", "updated_at": now}},
    )
    if not result.matched_count:
        changed = await db.followups.find_one(query, {"_id": 0})
        raise _notebook_conflict(NotebookConflictError(serialize_notebook_row(changed or before), ["referrer_id"]))
    await db.customers.update_one(
        {"id": before.get("customer_id"), "floor_id": floor_id}, {"$set": fields},
    )
    await log_event(
        event_type="project_followup.referrer_assigned", entity_type="followup", entity_id=row_id,
        actor=user, customer_id=before.get("customer_id"), floor_id=floor_id,
        payload={"referrer_id": fields["referrer_id"], "referrer_type": fields["referrer_type"]},
        summary=(f"Assigned {fields['referrer_name']}" if fields["referrer_name"] else "Cleared referrer"),
    )
    updated = await db.followups.find_one(query, {"_id": 0})
    return serialize_notebook_row(updated)


@router.post("/notebook/{floor_id}/{row_id}/convert")
async def convert_notebook(
    floor_id: str, row_id: str, body: NotebookConversionPayload,
    user: UserPublic = Depends(get_current_user),
):
    require_notebook_floor(floor_id, user)
    _reject_action_note(body.notes)
    before = await db.followups.find_one(notebook_query(user, floor_id, {"id": row_id}), {"_id": 0})
    if not before:
        raise HTTPException(status_code=404, detail="Notebook row not found")
    try:
        row = await convert_notebook_row(
            db, user=user, floor_id=floor_id, row_id=row_id,
            patch=body.model_dump(exclude={"updated_at", "notes"}, exclude_none=True), expected_updated_at=body.updated_at,
        )
    except NotebookConflictError as error:
        raise _notebook_conflict(error)
    except NotebookValidationError as error:
        raise HTTPException(status_code=422, detail=str(error))
    await log_event(
        event_type="project_followup.converted", entity_type="followup", entity_id=row_id, actor=user,
        customer_id=before.get("customer_id"), floor_id=floor_id,
        summary="Converted to quotation follow-up",
    )
    return row


@router.get("/notebook/{floor_id}/{row_id}/timeline")
async def notebook_timeline(floor_id: str, row_id: str, user: UserPublic = Depends(get_current_user)):
    require_notebook_floor(floor_id, user)
    row = await db.followups.find_one(notebook_query(user, floor_id, {"id": row_id}), {"_id": 0, "id": 1})
    if not row:
        raise HTTPException(status_code=404, detail="Notebook row not found")
    return await timeline_for(entity_type="followup", entity_id=row_id, limit=200, floor_ids=[floor_id])


@router.get("/config/customers")
async def eligible_lead_customers(q: str | None = None, limit: int = Query(250, ge=1, le=500), user: UserPublic = Depends(get_current_user)):
    query = {}
    if q:
        term = {"$regex": re.escape(q), "$options": "i"}
        query["$or"] = [{"name": term}, {"company": term}, {"phone": term}]
    return await db.customers.aggregate([
        {"$match": floor_query(user, query)}, {"$sort": {"name": 1, "id": 1}},
        {"$lookup": {"from": "quotations", "let": {"customer": "$id", "floor": "$floor_id"}, "pipeline": [
            {"$match": {"status": "ordered", "$expr": {"$and": [{"$eq": ["$customer_id", "$$customer"]}, {"$eq": ["$floor_id", "$$floor"]}]}}}, {"$limit": 1}, {"$project": {"_id": 0, "id": 1}},
        ], "as": "_orders"}}, {"$match": {"_orders": {"$size": 0}}}, {"$limit": limit},
        {"$project": {"_id": 0, "id": 1, "name": 1, "company": 1, "phone": 1, "tier": 1, "floor_id": 1}},
    ]).to_list(limit)


def _assigned_filter(user, value):
    return {"$or": [{"assigned_to": predicate(value)}, {"assignments": {"$elemMatch": {"user_id": predicate(value), "floor_id": floor_for_write(user)}}}]}


@router.get("/assignees")
async def lead_assignees(floor_ids: str | None = None, view: str = "walkin", source_id: str | None = None, customer_id: str | None = None, customer_phone: str | None = None, user: UserPublic = Depends(get_current_user)):
    if customer_phone and len(customer_phone) > 40:
        raise HTTPException(422, 'Phone must be at most 40 characters')
    floors = values(floor_ids) if floor_ids else [floor_for_write(user)]
    permitted = accessible_floor_ids(user)
    if permitted is not None:
        floors = [floor for floor in floors if floor in permitted]
        if not floors:
            return []
    floors, _ = await lead_scope.validate_scope(db, user, floors, [])
    stages = view == "quotation"
    if source_id:
        source = await db.followups.find_one(floor_query(user, {'id': source_id}), {'_id': 0})
        if not source:
            raise HTTPException(404, 'Follow-up not found')
        stages = await lead_scope.notebook_stages(db, source, floors)
    elif customer_id:
        customer = await db.customers.find_one(floor_query(user, {'id': customer_id}), {'_id': 0, 'phone': 1, 'floor_id': 1})
        if not customer:
            raise HTTPException(404, 'Customer not found')
        stages = await lead_scope.notebook_stages(db, {'floor_id': customer['floor_id'], 'customer_phone': customer.get('phone')}, floors)
    elif customer_phone:
        stages = await lead_scope.notebook_stages(db, {'floor_id': floor_for_write(user), 'customer_phone': customer_phone}, floors)
    return [{"id": staff["id"], "full_name": staff["full_name"], "role": staff["role"], "floor_ids": staff["eligible_floor_ids"]}
            for staff in await lead_scope.eligible_staff(db, floors, notebook_converted=stages) if staff["eligible_floor_ids"]]


@router.get("/team-assign")
async def team_assign(floor_ids: str | None = None, page: int = Query(1, ge=1), page_size: int = Query(50, ge=1, le=100), user: UserPublic = Depends(require_roles("owner"))):
    floors = values(floor_ids) if floor_ids else None
    query = {"category": {"$nin": list(lead_scope.POST_ORDER_CATEGORIES)}}
    if floors:
        floors, _ = await lead_scope.validate_scope(db, user, floors, [])
        query["floor_id"] = {"$in": floors}
    owner_rows = await db.users.find({"role": "owner"}, {"_id": 0, "id": 1}).to_list(1000)
    owner_ids = {row["id"] for row in owner_rows}
    # Page lead details, while aggregate counters cover the full selection.
    assignments = {"$filter": {"input": {"$ifNull": ["$assignments", []]}, "as": "a", "cond": {"$eq": ["$$a.floor_id", "$floor_id"]}}}
    assignment_read = db.followups.aggregate([
        {"$match": query},
        {"$set": {"_owners": assignments, "_lead": {"$ifNull": ["$lead_group_id", {"$ifNull": ["$walkin_id", {"$ifNull": ["$customer_id", "$id"]}]}]}}},
        {"$set": {"_owners": {"$cond": [{"$gt": [{"$size": "$_owners"}, 0]}, "$_owners", [{"user_id": "$assigned_to", "user_name": "$assigned_to_name"}]]}}},
        {"$unwind": "$_owners"}, {"$match": {"_owners.user_id": {"$type": "string", "$ne": ""}}},
        {"$group": {"_id": {"user": "$_owners.user_id", "lead": "$_lead"}, "user_name": {"$first": "$_owners.user_name"}, "completed": {"$min": {"$cond": [{"$in": ["$status", ["done", "dismissed"]]}, 1, 0]}}}},
        {"$group": {"_id": "$_id.user", "user_name": {"$first": "$user_name"}, "assigned_leads": {"$sum": 1}, "completed_leads": {"$sum": "$completed"}}},
    ]).to_list(1000)
    event_query = {"entity_type": "followup", "event_type": {"$in": ["followup.contacted", "followup.call_logged"]}}
    if floors:
        event_query["floor_id"] = {"$in": floors}
    contact_read = db.activity_events.aggregate([
        {"$match": event_query},
        {"$lookup": {"from": "followups", "localField": "entity_id", "foreignField": "id", "pipeline": [{"$match": query}, {"$project": {"_id": 0, "id": 1}}], "as": "_lead"}},
        {"$match": {"_lead.0": {"$exists": True}}},
        {"$group": {"_id": {"actor": "$actor_id", "lead": "$entity_id", "attempt": {"$ifNull": ["$payload.attempt_id", "$id"]}}, "is_call": {"$max": {"$cond": [{"$or": [{"$eq": ["$event_type", "followup.call_logged"]}, {"$eq": ["$payload.channel", "call"]}]}, 1, 0]}}}},
        {"$group": {"_id": "$_id.actor", "contacts_made": {"$sum": 1}, "calls_made": {"$sum": "$is_call"}}},
    ]).to_list(1000)
    rows, assignment_rows, contact_rows = await asyncio.gather(
        db.followups.find(query, {"_id": 0}).sort([("updated_at", -1), ("id", -1)]).skip((page-1)*page_size).limit(page_size+1).to_list(page_size+1),
        assignment_read, contact_read,
    )
    has_more = len(rows) > page_size
    rows = rows[:page_size]
    events = await db.activity_events.find({"entity_type": "followup", "entity_id": {"$in": [r["id"] for r in rows]}}, {"_id": 0}).sort("created_at", -1).limit(5000).to_list(5000)
    page_workers = {w["user_id"]: w for w in lead_scope.team_report(rows, events)}
    contacts = {r["_id"]: r for r in contact_rows}
    workers = []
    for summary in assignment_rows:
        uid = summary["_id"]
        if uid in owner_ids:
            continue
        worker = page_workers.get(uid, {"user_id": uid, "user_name": summary.get("user_name") or uid, "leads": []})
        worker.update({k: summary[k] for k in ("assigned_leads", "completed_leads")})
        worker.update({k: contacts.get(uid, {}).get(k, 0) for k in ("contacts_made", "calls_made")})
        workers.append(worker)
    workers.sort(key=lambda w: w["user_name"].casefold())
    return {"workers": workers, "has_more": has_more, "page": page, "page_size": page_size, "activity_truncated": len(events) == 5000}


@router.put("/{followup_id}/interests")
async def set_lead_interests(followup_id: str, body: FollowupInterestsPayload, user: UserPublic = Depends(get_current_user)):
    source = await db.followups.find_one(floor_query(user, {"id": followup_id}), {"_id": 0})
    if not source:
        raise HTTPException(404, "Follow-up not found")
    converted = getattr(user, "_custom_notebook_converted", None)
    if converted is not None and bool(source.get("is_converted")) != converted:
        raise HTTPException(403, "This notebook page is outside your assigned access")
    metadata = {}
    if "reference_source" in body.model_fields_set:
        metadata["reference_source"] = body.reference_source
    if "referrer_id" in body.model_fields_set:
        metadata.update(await lead_scope.referrer_fields(db, body.referrer_id, source["floor_id"], user))
    row = await lead_scope.update_interests(db, user=user, source=source, floor_ids=body.interested_floor_ids, assignments=body.assignments, metadata=metadata)
    await log_event(event_type="followup.assigned", entity_type="followup", entity_id=source["id"], actor=user, floor_id=source["floor_id"], customer_id=source.get("customer_id"), payload={"interested_floor_ids": row["interested_floor_ids"], "assignments": row["assignments"]}, summary="Updated interested floors and team assignment")
    return row


_ASSIGNMENT_STATUS_RANK = {"open": 0, "snoozed": 1, "done": 2, "dismissed": 2}


def _assignment_row(f: dict) -> dict:
    return {
        "id": f["id"], "assigned_to": f.get("assigned_to"), "assigned_to_name": f.get("assigned_to_name"),
        "customer_name": f.get("customer_name"), "reason": f.get("reason"), "category": f.get("category"),
        "status": f.get("status"), "bucket": compute_bucket(f),
        "days_pending": age_days(parse_iso(f.get("created_at"))),
        "due_at": f.get("due_at"), "created_at": f.get("created_at"),
    }


def _assignment_sort_key(row: dict) -> tuple:
    return (_ASSIGNMENT_STATUS_RANK.get(row["status"], 3), -row["days_pending"])


@router.get("/assignments")
async def list_assignments(
    include_completed: bool = False,
    user: UserPublic = Depends(require_min_role("manager")),
):
    """Who has what assigned, how long it's been pending, and whether it's
    done — manager/admin/owner only. See
    docs/superpowers/specs/2026-07-27-followups-revamp-design.md."""
    status_filter = ["open", "snoozed"] + (["done", "dismissed"] if include_completed else [])
    docs = await db.followups.find(
        floor_query(user, {"assigned_to": {"$ne": None}, "status": {"$in": status_filter}}),
        {"_id": 0},
    ).to_list(5000)
    owners = await db.users.find({"role": "owner"}, {"_id": 0, "id": 1}).to_list(1000)
    owner_ids = {row["id"] for row in owners}
    rows = [_assignment_row(f) for f in docs if f.get("assigned_to") not in owner_ids]
    rows.sort(key=_assignment_sort_key)
    return rows


# ─────────────────────────────────────────────────────────────────────────────
# Detail — powers the Customer Context Panel
# ─────────────────────────────────────────────────────────────────────────────
@router.get("/{followup_id}")
async def get_detail(followup_id: str, user: UserPublic = Depends(get_current_user)):
    f = await _selected_followup_or_404(followup_id, user)
    if f.get("category") in lead_scope.POST_ORDER_CATEGORIES or f.get("rule_type") in lead_scope.POST_ORDER_RULES or f.get("purchase_id"):
        raise HTTPException(404, "This activity belongs in Orders")
    floor_id = f["floor_id"]
    f["bucket"] = compute_bucket(f)
    f["effective_priority_level"] = f.get("manual_priority_override") or f.get("priority_level")
    customer, pending_quotations, timeline = await asyncio.gather(
        db.customers.find_one({"id": f["customer_id"], "floor_id": floor_id}, {"_id": 0, "password_hash": 0}),
        db.quotations.find({"customer_id": f["customer_id"], "floor_id": floor_id,
                            "status": {"$in": ["draft", "sent", "pending_approval", "approved", "selection"]}},
                           {"_id": 0, "id": 1, "number": 1, "status": 1, "grand_total": 1,
                            "valid_until": 1, "updated_at": 1, "created_by_name": 1}).sort("updated_at", -1).to_list(200),
        timeline_for(entity_type="followup", entity_id=followup_id, limit=60, floor_ids=[floor_id]),
    )
    await _enrich_followup_references([f], user, [customer])
    creators = Counter(q.get("created_by_name") for q in pending_quotations if q.get("created_by_name"))
    preferred = creators.most_common(1)[0][0] if creators and floor_id != "third-floor" else None
    touches = [parse_iso(q.get("updated_at")) for q in pending_quotations]
    touches = [t for t in touches if t]
    days_silent = age_days(max(touches)) if touches else 0
    return {
        "followup": f, "customer": customer or {},
        "stats": {"lifetime_revenue": 0, "outstanding_total": 0, "pending_orders": 0,
                  "conversion_rate": 0, "average_order_value": 0,
                  "pending_quotations": len(pending_quotations), "preferred_salesperson": preferred,
                  "risk_level": "medium" if days_silent >= 14 else "low"},
        "quotations": pending_quotations[:12], "payments": [], "purchases": [], "timeline": timeline,
    }


# ─────────────────────────────────────────────────────────────────────────────
# Manual create + edit
# ─────────────────────────────────────────────────────────────────────────────
def _required_create_assignments(floors, assignments, assigned_to=None, source_floor=None):
    """Every newly created floor-local lead needs an explicit worker selection."""
    selected = list(assignments or [])
    if assigned_to and not any((a.floor_id if hasattr(a, "floor_id") else a.get("floor_id")) == source_floor for a in selected):
        selected = [{"floor_id": source_floor, "user_id": assigned_to}]
    assigned_floors = {(a.floor_id if hasattr(a, "floor_id") else a.get("floor_id")) for a in selected
                       if (a.user_id if hasattr(a, "user_id") else a.get("user_id"))}
    if not floors or any(floor not in assigned_floors for floor in floors):
        raise HTTPException(422, "Select a team member for each interested floor")
    return selected


@router.post("", response_model=Followup)
async def create_followup(body: FollowupCreate, user: UserPublic = Depends(get_current_user)):
    _reject_action_note(body.notes)
    floor_id = floor_for_write(user)
    request_id = str(uuid5(NAMESPACE_URL, f'followup-create:{user.id}:{floor_id}:{body.request_key}')) if body.request_key else None
    fingerprint = hashlib.sha256(json.dumps(body.model_dump(mode='json', exclude={'request_key'}), sort_keys=True, separators=(',', ':')).encode()).hexdigest()
    previous = await db.followups.find_one({'id': request_id, 'floor_id': floor_id}, {'_id': 0}) if request_id else None
    if previous:
        if previous.get('create_request_fingerprint') != fingerprint:
            raise HTTPException(409, 'This request key was already used for a different follow-up; start a new reminder')
        if previous.get('create_request_complete'):
            return Followup(**previous)
    if body.category in lead_scope.POST_ORDER_CATEGORIES or body.purchase_id:
        raise HTTPException(422, "Follow-ups cover only activities before an order is placed")
    if body.category == "sales":
        raise HTTPException(422, "Choose General, Quotation, Selection or Walk-in as the follow-up type")
    interested_floors = body.interested_floor_ids or [floor_id]
    selected_assignments = _required_create_assignments(interested_floors, body.assignments, body.assigned_to, floor_id)
    new_customer_details = None
    if body.customer_id and (body.new_customer_name or body.new_customer_phone):
        raise HTTPException(422, "Choose an existing customer or enter a new customer's details")
    if body.customer_id:
        cust = await db.customers.find_one(floor_query(user, {"id": body.customer_id}), {"_id": 0})
    elif body.new_customer_name and body.new_customer_phone:
        mobile = normalize_mobile(body.new_customer_phone)
        if not body.new_customer_name.strip():
            raise HTTPException(422, "Customer name is required")
        if len(mobile) != 10:
            raise HTTPException(422, "Mobile number must contain 10 digits")
        new_customer_details = (body.new_customer_name.strip(), body.new_customer_phone.strip())
        cust = {
            "id": str(uuid5(NAMESPACE_URL, f'followup-customer:{request_id}')) if request_id else str(uuid4()),
            "name": new_customer_details[0], "phone": new_customer_details[1], "tier": "retail",
        }
    else:
        raise HTTPException(422, "Choose a customer or add a customer name and mobile number")
    if not cust:
        raise HTTPException(status_code=404, detail="Customer not found")
    quotation = await db.quotations.find_one(floor_query(user, {"id": body.quotation_id}), {"_id": 0}) if body.quotation_id else None
    if body.quotation_id and not quotation:
        raise HTTPException(404, "Quotation not found")
    if quotation and quotation.get("status") == "ordered":
        raise HTTPException(409, "This quotation is already an order")
    if not new_customer_details and await db.quotations.find_one({"customer_id": cust["id"], "floor_id": floor_id, "status": "ordered"}, {"_id": 0, "id": 1}):
        raise HTTPException(409, "This customer already placed an order")
    tier = cust.get("tier", "retail")
    value = float(quotation.get("grand_total", 0)) if quotation else 0.0
    score, level = score_followup(value, 0, 12, tier)

    assigned_to_name = None

    action_label = {"call": "Call customer", "whatsapp": "Send WhatsApp", "email": "Send Email", "visit": "Schedule showroom visit"}
    f = Followup(
        **({'id': request_id} if request_id else {}),
        source_key=None, rule_type="manual", category=body.category,
        customer_id=cust["id"], customer_name=cust.get("name"),
        customer_phone=cust.get("phone"), customer_tier=tier,
        quotation_id=quotation.get("id") if quotation else None,
        quotation_number=quotation.get("number") if quotation else None,
        purchase_id=body.purchase_id,
        value=value, reason=body.reason,
        reason_factors=reason_factors_for(value, 0, "Manual reminder", tier),
        next_action=action_label.get(body.channel, "Call customer"),
        next_action_reason=body.reason,
        suggested_channel=body.channel, priority_score=score,
        priority_level=body.priority_level or level,
        due_at=body.due_at or now_iso(), is_automated=False,
        assigned_to=None, assigned_to_name=assigned_to_name,
        floor_id=floor_id,
    )
    if previous:
        f = Followup(**previous)
    metadata = {"reference_source": body.reference_source}
    if body.referrer_id:
        metadata.update(await lead_scope.referrer_fields(db, body.referrer_id, f.floor_id, user))
    if selected_assignments:
        await lead_scope.update_interests(db, user=user, source=f.model_dump(), floor_ids=interested_floors, assignments=selected_assignments, metadata=metadata, dry_run=True)
    if new_customer_details:
        try:
            cust, customer_created = await resolve_or_create_customer(
                db, user=user, floor_id=floor_id,
                name=new_customer_details[0], phone=new_customer_details[1], address=None,
                return_created=True,
            )
        except NotebookValidationError as exc:
            raise HTTPException(422, str(exc)) from exc
        f = f.model_copy(update={
            "customer_id": cust["id"],
            "customer_name": cust.get("name"),
            "customer_phone": cust.get("phone"),
            "customer_tier": cust.get("tier", "retail"),
        })
        if await db.quotations.find_one({"customer_id": cust["id"], "floor_id": floor_id, "status": "ordered"}, {"_id": 0, "id": 1}):
            raise HTTPException(409, "This customer already placed an order")
        if customer_created:
            await log_event(
                event_type="customer.created", entity_type="customer", entity_id=cust["id"],
                customer_id=cust["id"], actor=user, summary="Customer Created with Follow-up",
            )
    if request_id:
        identity = {'_id': f'followup-create:{request_id}'}
        await db.followups.update_one(identity, {'$setOnInsert': {**f.model_dump(), 'create_request_fingerprint': fingerprint}}, upsert=True)
        stored = await db.followups.find_one(identity, {'_id': 0})
        if stored.get('create_request_fingerprint') != fingerprint:
            raise HTTPException(409, 'This request key was already used for a different follow-up; start a new reminder')
        if stored.get('create_request_complete'):
            return Followup(**stored)
        f = Followup(**stored)
    else:
        await db.followups.insert_one(f.model_dump())
    if selected_assignments:
        document = await lead_scope.update_interests(db, user=user, source=f.model_dump(), floor_ids=interested_floors, assignments=selected_assignments, metadata=metadata)
        f = Followup(**document)
    completion = await db.followups.update_one({'id': f.id, 'create_request_complete': {'$ne': True}}, {'$set': {'create_request_complete': True}}) if request_id else None
    if completion is None or completion.matched_count:
        await log_event(
            event_type="followup.created", entity_type="followup", entity_id=f.id, actor=user,
            customer_id=f.customer_id, quotation_id=f.quotation_id,
            summary=f"Manual follow-up created — {f.reason}",
        )
    return f


@router.patch("/{followup_id}")
async def update_followup(followup_id: str, body: FollowupUpdate, user: UserPublic = Depends(get_current_user)):
    _reject_action_note(body.notes)
    f = await db.followups.find_one(floor_query(user, {"id": followup_id}), {"_id": 0})
    if not f:
        raise HTTPException(status_code=404, detail="Follow-up not found")
    patch = body.model_dump(exclude_unset=True)
    patch.pop("notes", None)
    scope_keys = {"interested_floor_ids", "assignments", "reference_source", "referrer_id"}
    scope_patch = {k: patch.pop(k) for k in list(patch) if k in scope_keys}
    if scope_patch:
        metadata = {k: v for k, v in scope_patch.items() if k == "reference_source"}
        if "referrer_id" in scope_patch:
            metadata.update(await lead_scope.referrer_fields(db, scope_patch["referrer_id"], f["floor_id"], user))
        f = await lead_scope.update_interests(db, user=user, source=f, floor_ids=scope_patch.get("interested_floor_ids") or f.get("interested_floor_ids") or [f["floor_id"]], assignments=scope_patch.get("assignments"), metadata=metadata)
    if "assigned_to" in patch and patch["assigned_to"]:
        await lead_scope.validate_scope(db, user, [f["floor_id"]], [{"floor_id": f["floor_id"], "user_id": patch["assigned_to"]}])
        au = await db.users.find_one({"id": patch["assigned_to"]}, {"_id": 0, "full_name": 1})
        patch["assigned_to_name"] = au["full_name"] if au else None
        await log_event(
            event_type="followup.assigned", entity_type="followup", entity_id=followup_id, actor=user,
            customer_id=f.get("customer_id"),
            summary=f"Assigned to {patch.get('assigned_to_name') or '—'}",
        )
    if patch.get("status") == "dismissed":
        patch["completed_at"] = now_iso()
        patch["resolution_note"] = "Follow-up dismissed"
        await log_event(
            event_type="followup.dismissed", entity_type="followup", entity_id=followup_id, actor=user,
            customer_id=f.get("customer_id"),
            summary=f"Follow-up dismissed — {f.get('reason')}",
        )
    patch["updated_at"] = now_iso()
    await db.followups.update_one(floor_query(user, {"id": followup_id}), {"$set": patch})
    return await db.followups.find_one(floor_query(user, {"id": followup_id}), {"_id": 0})


@router.delete("/{followup_id}")
async def delete_followup(
    followup_id: str,
    user: UserPublic = Depends(require_min_role("manager")),
    notes: str | None = Query(None),
):
    """Permanently remove a follow-up during manager data cleanup."""
    _reject_action_note(notes)
    f = await _selected_followup_or_404(followup_id, user)
    result = await db.followups.delete_one(floor_query(user, {"id": followup_id}))
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Follow-up not found")
    activity = await db.activity_events.delete_many(floor_query(user, {
        "$or": [
            {"entity_type": "followup", "entity_id": followup_id},
            {"followup_id": followup_id},
        ],
    }))
    await log_event(
        event_type="followup.deleted", entity_type="followup", entity_id=followup_id, actor=user,
        customer_id=f.get("customer_id"), floor_id=f.get("floor_id"),
        summary="Follow-up deleted",
    )
    return {"ok": True, "followup_id": followup_id, "deleted": {"followups": result.deleted_count, "activity_events": activity.deleted_count}}


# ─────────────────────────────────────────────────────────────────────────────
# Actions — snooze / complete / contact / call outcome
# ─────────────────────────────────────────────────────────────────────────────
@router.post("/{followup_id}/snooze")
async def snooze_followup(followup_id: str, body: FollowupSnoozePayload, user: UserPublic = Depends(get_current_user)):
    _reject_action_note(body.notes)
    f = await db.followups.find_one(floor_query(user, {"id": followup_id}), {"_id": 0})
    if not f:
        raise HTTPException(status_code=404, detail="Follow-up not found")
    now = datetime.now(timezone.utc)
    if body.until:
        until = parse_iso(body.until) or (now + timedelta(hours=1))
    elif body.minutes:
        until = now + timedelta(minutes=body.minutes)
    elif body.preset == "15m":
        until = now + timedelta(minutes=15)
    elif body.preset == "1h":
        until = now + timedelta(hours=1)
    elif body.preset == "tomorrow":
        until = (now + timedelta(days=1)).replace(hour=9, minute=0, second=0, microsecond=0)
    elif body.preset == "next_week":
        until = (now + timedelta(days=7)).replace(hour=9, minute=0, second=0, microsecond=0)
    else:
        until = now + timedelta(hours=1)

    await db.followups.update_one(floor_query(user, {"id": followup_id}), {"$set": {
        "status": "snoozed", "snoozed_until": until.isoformat(), "updated_at": now_iso(),
    }})
    await log_event(
        event_type="followup.snoozed", entity_type="followup", entity_id=followup_id, actor=user,
        customer_id=f.get("customer_id"),
        summary=f"Snoozed until {until.strftime('%d %b, %I:%M %p')}",
    )
    return await db.followups.find_one(floor_query(user, {"id": followup_id}), {"_id": 0})


@router.post("/{followup_id}/complete")
async def complete_followup(followup_id: str, body: FollowupCompletePayload, user: UserPublic = Depends(get_current_user)):
    _reject_action_note(body.notes)
    f = await db.followups.find_one(floor_query(user, {"id": followup_id}), {"_id": 0})
    if not f:
        raise HTTPException(status_code=404, detail="Follow-up not found")
    await db.followups.update_one(floor_query(user, {"id": followup_id}), {"$set": {
        "status": "done", "completed_at": now_iso(),
        "updated_at": now_iso(),
    }})
    await log_event(
        event_type="followup.completed", entity_type="followup", entity_id=followup_id, actor=user,
        customer_id=f.get("customer_id"), quotation_id=f.get("quotation_id"), purchase_id=f.get("purchase_id"),
        summary="Follow-up marked complete",
    )
    return await db.followups.find_one(floor_query(user, {"id": followup_id}), {"_id": 0})


@router.post("/{followup_id}/contact")
async def contact_followup(followup_id: str, body: FollowupContactPayload, user: UserPublic = Depends(get_current_user)):
    _reject_action_note(body.notes)
    from routes.payment_routes import _clean_phone
    from services.email_service import build_email
    from services.messaging_service import build_message

    identity = {"id": followup_id}
    converted = getattr(user, "_custom_notebook_converted", None)
    if converted is not None:
        identity.update(is_converted=converted, notebook_key={"$type": "string"})
    contact_query = floor_query(user, identity)
    f = await db.followups.find_one(contact_query, {"_id": 0})
    if not f:
        raise HTTPException(status_code=404, detail="Follow-up not found")
    now = now_iso()
    attempt_id = str(uuid4())
    contact_patch = {"last_contacted_at": now, "updated_at": now}
    increments = {"contact_attempts": 1}
    if body.channel == "call":
        contact_patch.update(pending_call_id=attempt_id, pending_call_actor_id=user.id)
        actor_key = hashlib.sha256(user.id.encode()).hexdigest()
        contact_patch[f"pending_call_ids.{actor_key}"] = attempt_id
        increments["call_attempts"] = 1
    result = await db.followups.update_one(contact_query, {"$set": contact_patch, "$inc": increments})
    if converted is not None and not result.matched_count:
        raise HTTPException(status_code=409, detail="This lead moved to another notebook page; reopen that page")
    await log_event(
        event_type="followup.contacted", entity_type="followup", entity_id=followup_id, actor=user,
        customer_id=f.get("customer_id"), quotation_id=f.get("quotation_id"), purchase_id=f.get("purchase_id"),
        payload={"channel": body.channel, "attempt_id": attempt_id},
        summary=f"{body.channel.title()} — {f.get('customer_name')}",
    )
    phone = _clean_phone(f.get("customer_phone"))
    result: dict = {"channel": body.channel, "phone": phone}
    # Category-aware message templates (services/messaging_service.py,
    # services/email_service.py) — provider-based, deep-link only in Phase 1.
    ctx = {
        "customer_name": (f.get("customer_name") or "there").split()[0],
        "quotation_number": f.get("quotation_number") or "",
        "outstanding_amount": money_short(f.get("value", 0)),
        "salesperson_name": f.get("assigned_to_name") or "",
        "reason": f.get("reason") or "Just checking in!",
    }
    if body.channel == "whatsapp":
        msg = build_message(f.get("category", "general"), phone, ctx)
        result["message"] = msg["message"]
        result["wa_url"] = msg["url"]
    elif body.channel == "email":
        cust = await db.customers.find_one(
            {"id": f["customer_id"], "floor_id": f.get("floor_id")}, {"_id": 0, "email": 1},
        )
        result["email"] = cust.get("email") if cust else None
        template_key = "payment_reminder" if f.get("category") == "payment" else "quotation"
        email = build_email(template_key, result["email"], {**ctx, "invoice_number": f.get("quotation_number") or "", "order_number": f.get("quotation_number") or ""})
        result["subject"] = email["subject"]
        result["body"] = email["body"]
        result["mailto_url"] = email["mailto_url"]
    return result


async def _record_lost_quotation_followup(followup: dict, update: dict, user: UserPublic) -> None:
    """Commit a linked sales loss with its quotation so Place Order cannot race it."""
    from db import client

    floor = followup["floor_id"]
    require_floor_access(floor, user)
    query = {"id": followup["id"], "floor_id": floor}
    async def record(session):
        current = await db.followups.find_one(query, {"_id": 0}, session=session)
        if not current:
            raise HTTPException(404, "Follow-up not found")
        if current.get("status") not in {"open", "snoozed"} or current.get("updated_at") != followup.get("updated_at") or current.get("quotation_id") != followup.get("quotation_id"):
            raise HTTPException(409, "This follow-up changed. Reopen it before recording the loss.")
        quotation_query = {"id": followup["quotation_id"], "floor_id": floor}
        quotation = await db.quotations.find_one(quotation_query, {"_id": 0}, session=session)
        if not quotation or quotation.get("customer_id") != current.get("customer_id"):
            raise HTTPException(404, "Linked quotation not found on this floor")
        if quotation.get("status") == "ordered" or quotation.get("customer_outcome") == "won":
            raise HTTPException(409, "This quotation already has a placed order or won outcome and cannot be marked lost")
        patch = update["$set"]
        await db.quotations.update_one(quotation_query, {"$set": {
            "customer_outcome": "lost", "outcome_reason": patch["lost_reason"], "updated_at": patch["updated_at"],
        }}, session=session)
        await db.followups.update_one(query, update, session=session)
        # Retire other reminders for this document while retaining their assignments/history.
        await db.followups.update_many({"quotation_id": quotation["id"], "floor_id": floor, "status": {"$in": ["open", "snoozed"]}}, {"$set": {
            key: patch[key] for key in ("status", "completed_at", "completed_outcome", "resolution_note", "lost_reason", "updated_at")
        }}, session=session)

    async with await client.start_session() as session:
        await session.with_transaction(record)


@router.post("/{followup_id}/log-call")
async def log_call(followup_id: str, body: FollowupCallOutcomePayload, user: UserPublic = Depends(get_current_user)):
    outcome = "lost" if body.outcome == "rejected" else body.outcome
    notes = _required_action_note(body.notes) if outcome == "lost" else None
    if outcome != "lost":
        _reject_action_note(body.notes)
    f = await db.followups.find_one(floor_query(user, {"id": followup_id}), {"_id": 0})
    if not f:
        raise HTTPException(status_code=404, detail="Follow-up not found")
    now_dt = datetime.now(timezone.utc)
    if outcome in {"won", "converted"}:
        from services.order_confirmation import placed_order
        order = await placed_order(db, f)
        if not order:
            raise HTTPException(status_code=409, detail="Won is recorded only after an order is placed")
    actor_key = hashlib.sha256(user.id.encode()).hexdigest()
    pending_attempt = (f.get("pending_call_ids") or {}).get(actor_key)
    if not pending_attempt and f.get("pending_call_actor_id") == user.id:
        pending_attempt = f.get("pending_call_id")
    attempt_id = pending_attempt or str(uuid4())
    patch: dict = {
        "last_contacted_at": now_dt.isoformat(), "updated_at": now_iso(),
        "call_attempts": (f.get("call_attempts") or 0) + (0 if pending_attempt else 1),
        "contact_attempts": (f.get("contact_attempts") or 0) + (0 if pending_attempt else 1),
        "latest_response": (body.response or "").strip() or outcome.replace("_", " "),
    }
    if outcome in {"won", "converted"}:
        patch["auto_resolved"] = True
        if order.get("created_by"):
            patch.update({"assigned_to": order["created_by"], "assigned_to_name": order.get("created_by_name")})
    if outcome == "lost":
        patch["notes"] = notes
    next_created = None

    if outcome == "lost" and not (body.notes or "").strip():
        raise HTTPException(status_code=422, detail="A reason is required when marking a client as lost")

    if outcome == "pending" and body.next_followup_at is None:
        raise HTTPException(status_code=422, detail="Choose the next follow-up date for a pending client")

    if outcome == "pending":
        next_due = body.next_followup_at
        if next_due is None or (next_due.replace(tzinfo=timezone.utc) if next_due.tzinfo is None else next_due.astimezone(timezone.utc)) <= now_dt:
            raise HTTPException(422, "The next follow-up date must be in the future")
    successor_assignments = None
    if outcome in ("interested", "call_back", "pending"):
        selected = body.assignments if body.assignments is not None else f.get("assignments")
        selected = _required_create_assignments([f["floor_id"]], selected, f.get("assigned_to") if body.assignments is None else None, f["floor_id"])
        local_selected = [item for item in selected if (item.floor_id if hasattr(item, "floor_id") else item["floor_id"]) == f["floor_id"]]
        successor_assignments = await lead_scope.validate_assignments(db, [f["floor_id"]], local_selected, bool(f.get("is_converted")))
    if body.interested_floor_ids or body.assignments is not None:
        f = await lead_scope.update_interests(db, user=user, source=f, floor_ids=body.interested_floor_ids or f.get("interested_floor_ids") or [f["floor_id"]], assignments=body.assignments)
    previous_assignee = f.get("assigned_to")
    if previous_assignee and await db.users.find_one({"id": previous_assignee, "role": "owner"}, {"_id": 0, "id": 1}):
        previous_assignee = None
    if successor_assignments:
        previous_assignee = successor_assignments[0]["user_id"]
        f = {**f, "assigned_to": previous_assignee, "assigned_to_name": successor_assignments[0]["user_name"],
             "assignments": [a for a in f.get("assignments", []) if a["floor_id"] != f["floor_id"]] + successor_assignments}
    if outcome == "interested":
        # A transferred lead stays with the recipient, including the source
        # row and the new reminder; interest is not a placed/won order.
        if previous_assignee or user.role != "owner":
            patch.update({"assigned_to": previous_assignee or user.id,
                          "assigned_to_name": f.get("assigned_to_name") if previous_assignee else user.full_name})
        else:
            patch.update({"assigned_to": None, "assigned_to_name": None})

    if outcome in ("interested", "call_back", "pending"):
        patch.update({"status": "done", "completed_at": now_iso(), "completed_outcome": outcome})
        if outcome == "pending":
            due = body.next_followup_at
            if due.tzinfo is None:
                due = due.replace(tzinfo=timezone.utc)
            else:
                due = due.astimezone(timezone.utc)
            if due <= now_dt:
                raise HTTPException(status_code=422, detail="The next follow-up date must be in the future")
        else:
            due = now_dt + timedelta(days=1 if outcome == "call_back" else 2)
        next_reason = {
            "call_back": "Call back requested",
            "interested": "Customer interested — follow up on their decision",
            "pending": "Customer asked to reconnect later",
        }[outcome]
        nf = Followup(
            rule_type="manual", category=f.get("category", "general"),
            origin=f.get("origin"), walkin_id=f.get("walkin_id"),
            customer_id=f["customer_id"], customer_name=f["customer_name"], customer_phone=f.get("customer_phone"),
            customer_tier=f.get("customer_tier", "retail"), quotation_id=f.get("quotation_id"),
            quotation_number=f.get("quotation_number"), purchase_id=f.get("purchase_id"),
            purchase_number=f.get("purchase_number"), value=f.get("value", 0),
            reason=next_reason,
            reason_factors=[f.get("reason", "")] if f.get("reason") else [],
            next_action="Call customer", next_action_reason="Scheduled automatically from the previous call outcome.",
            suggested_channel="call", priority_score=f.get("priority_score", 50), priority_level=f.get("priority_level", "medium"),
            due_at=due.isoformat(), is_automated=False,
            assigned_to=previous_assignee or (None if user.role == "owner" else user.id),
            assigned_to_name=(f.get("assigned_to_name") if previous_assignee else (None if user.role == "owner" else user.full_name)),
            tags=f.get("tags", []),
            floor_id=floor_inherit(f),
            lead_group_id=f.get("lead_group_id"), interested_floor_ids=f.get("interested_floor_ids", []),
            assignments=f.get("assignments", []), reference_source=f.get("reference_source"),
            referrer_id=f.get("referrer_id"), referrer_name=f.get("referrer_name"), referrer_type=f.get("referrer_type"),
        )
        await db.followups.insert_one(nf.model_dump())
        next_created = nf.id
    elif outcome == "no_answer":
        attempts = patch["contact_attempts"]
        patch["contact_attempts"] = attempts
        if attempts >= 2:
            # Escalate — stop same-day retries after the 2nd miss; push to
            # tomorrow morning and bump urgency so it doesn't get buried.
            next_due = (now_dt + timedelta(days=1)).replace(hour=9, minute=30, second=0, microsecond=0)
            bumped_score = min(100, (f.get("priority_score") or 0) + 10)
            bumped_level = "critical" if bumped_score >= 80 else "high" if bumped_score >= 60 else f.get("priority_level", "medium")
            patch["due_at"] = next_due.isoformat()
            patch["priority_score"] = bumped_score
            patch["priority_level"] = bumped_level
        else:
            patch["due_at"] = (now_dt + timedelta(hours=4)).isoformat()
    elif outcome == "rejected":
        patch.update({
            "status": "dismissed", "completed_at": now_iso(),
            "completed_outcome": "rejected", "resolution_note": "Customer rejected",
        })
    elif outcome == "converted":
        patch.update({
            "status": "done", "completed_at": now_iso(),
            "completed_outcome": "converted", "resolution_note": "Converted!",
        })
    elif outcome == "won":
        patch.update({
            "status": "done", "completed_at": now_iso(),
            "completed_outcome": "won", "resolution_note": "Client won",
        })
    elif outcome == "lost":
        patch.update({
            "status": "dismissed", "completed_at": now_iso(),
            "completed_outcome": "lost", "resolution_note": notes, "lost_reason": notes,
        })

    # Contact initiation already incremented the counters. An outcome must
    # not overwrite increments from another worker contacting this same lead.
    patch.pop("contact_attempts", None)
    patch.pop("call_attempts", None)
    update = {"$set": patch}
    if not pending_attempt:
        update["$inc"] = {"contact_attempts": 1, "call_attempts": 1}
    if outcome == "lost" and f.get("quotation_id") and f.get("floor_id") in {"ground-floor", "first-floor"}:
        await _record_lost_quotation_followup(f, update, user)
    else:
        await db.followups.update_one(floor_query(user, {"id": followup_id}), update)
    if pending_attempt:
        # A second initiation by this worker may finish while this outcome is
        # awaiting scope writes. Only consume the attempt that we actually read.
        await db.followups.update_one(floor_query(user, {'id': followup_id, f'pending_call_ids.{actor_key}': pending_attempt}), {'$unset': {f'pending_call_ids.{actor_key}': ''}})
        await db.followups.update_one(floor_query(user, {'id': followup_id, 'pending_call_id': pending_attempt, 'pending_call_actor_id': user.id}), {'$set': {'pending_call_id': None, 'pending_call_actor_id': None}})
    await log_event(
        event_type="followup.call_logged", entity_type="followup", entity_id=followup_id, actor=user,
        customer_id=f.get("customer_id"), quotation_id=f.get("quotation_id"), purchase_id=f.get("purchase_id"),
        payload={"outcome": outcome, "next_followup_id": next_created, "notes": notes, "response": body.response, "attempt_id": attempt_id},
        summary=f"Call logged — {outcome.replace('_', ' ').title()}" + (f" — {notes}" if notes else ""),
    )
    return await db.followups.find_one(floor_query(user, {"id": followup_id}), {"_id": 0})
